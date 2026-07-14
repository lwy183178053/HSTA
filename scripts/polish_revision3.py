from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image


REWRITES = {
    "面向轻量级加密 QUIC/TLS 流量分类的混合状态空间转换与注意力建模方法": (
        "HSTA：面向轻量级加密 QUIC/TLS 流量分类的混合状态空间转换与注意力模型"
    ),
    "Hybrid State-Space Transformation and Attention Modeling for Lightweight Encrypted QUIC/TLS Traffic Classification": (
        "HSTA: Hybrid State-Space Transformation and Attention for Lightweight Encrypted QUIC/TLS Traffic Classification"
    ),
    "随着 TLS 1.3 与 QUIC 等加密通信协议广泛部署": (
        "随着 TLS 1.3 与 QUIC 等现代加密协议广泛部署，依赖端口号、明文载荷或深度包检测的传统流量识别方式已难以满足细粒度网络管理需求。为在明文不可见条件下充分挖掘包级行为信息，本文提出混合状态空间转换与注意力模型（Hybrid State-Space Transformation and Attention, HSTA）。HSTA 以连续 Mamba 层建模包大小、方向和到达时间间隔形成的动态依赖，通过 Transition MLP 将序列状态转换到高层判别空间，并利用后置注意力与 Refinement MLP 完成关键模式重加权和特征重整。本文基于 CESNET-QUIC22 和 CESNET-TLS22 构建 QUIC-40、QUIC-60、TLS-40 和 TLS-60 四个任务，在统一轻量包级输入和一致训练协议下开展系统评估。HSTA 在四个任务上的 Macro-F1 分别达到 91.09±0.20%、90.28±0.39%、96.78±0.14% 和 96.29±0.17%，均取得最高平均性能；在更具挑战性的 QUIC-40 和 QUIC-60 上，相较 NetMamba 分别提升 1.21 和 1.32 个百分点，相较 Transformer 分别提升 2.52 和 3.13 个百分点。注意力位置消融进一步表明，完整 HSTA 在四个任务上相较最强变体提升 0.41–1.91 个百分点，验证了“状态空间编码—表示转换—高层重加权—特征重整”的结构协同。效率结果显示，HSTA 相较 Transformer 减少约 41.5% 的参数量和 65.9% 的单样本 GFLOPs，并在 batch size 为 1 和 32 时保持更低平均延迟。跨协议 full fine-tuning 在 QUIC→TLS-40 和 QUIC→TLS-60 上进一步达到 96.96% 和 96.60% Macro-F1。上述结果表明，HSTA 在分类精度、结构有效性、在线推理效率和协议适配能力之间形成了稳定而有竞争力的综合优势。"
    ),
    "早期流量分类通常依赖端口号匹配、深度包检测或人工统计特征": (
        "早期流量分类通常依赖端口号匹配、深度包检测或人工统计特征，这类方法需要专家根据包间隔、包大小、协议字段和流持续时间等指标进行特征工程[4]。随着动态端口、应用伪装、协议升级和端到端加密普及，研究重点逐步转向从加密可见行为中自动学习判别表示[5]。Deep Packet 展示了深度模型直接学习加密流量表示的能力[11]，ET-BERT 通过预训练 Transformer 建模报文上下文[16]，NetMamba 将选择性状态空间模型用于高效流量序列建模[14]，TrafficFormer 则进一步推进面向流量数据的预训练表征学习[27]。面向现代 QUIC/TLS 的细粒度分类，模型需要同时解决三个关键问题：在轻量侧信道输入中提取充分信息，在行为相似类别之间形成清晰边界，以及在不同协议分布之间保持可适配性。实际网络监测还要求模型兼顾分类精度、推理延迟和计算开销，这为结构设计提出了更完整的综合目标。"
    ),
    "针对上述问题，细粒度 QUIC/TLS 加密流量分类不仅需要捕获": (
        "针对上述需求，本文提出混合状态空间转换与注意力模型（Hybrid State-Space Transformation and Attention, HSTA），将高效序列建模与判别模式选择组织为连续的特征形成过程。连续 Mamba 层首先捕获方向切换、包大小变化、到达时间间隔和突发传输形成的包级动态依赖；Transition MLP 随后完成高层表示转换；注意力模块在已编码的高层空间中突出关键包段和特征组合；Refinement MLP 最终重整增强后的序列表示。该设计使模型能够在仅使用三维包级侧信道特征的条件下，形成兼具顺序感知能力和全局判别能力的紧凑流量表示。"
    ),
    "1. 面向现代 QUIC/TLS 加密流量分类中的实际问题": (
        "1. 构建覆盖 QUIC 与 TLS、40 类与 60 类设置的四个细粒度任务，并在统一 DataZoo 划分、固定 30 包序列和三维侧信道输入下建立可重复的评估基准。该设置同时覆盖协议差异、类别规模变化和类别不均衡，为检验现代加密流量模型的综合能力提供了充分实验基础。"
    ),
    "2. 针对侧信道特征信息有限、细粒度类别行为相似的问题": (
        "2. 提出以 HSTA 模块为核心的轻量分类模型，将双层 Mamba 状态空间编码、Transition MLP 表示转换、后置注意力重加权和 Refinement MLP 特征重整进行协同设计。该结构能够从紧凑包级输入中连续提取动态依赖与关键判别模式，并以约 0.604M 参数实现四个任务上的最高平均 Macro-F1。"
    ),
    "3. 围绕关键模式选择是否真正有效的问题": (
        "3. 通过无注意力、前置注意力、中置注意力和完整 HSTA 四种结构的三随机种子消融，系统验证注意力位置与表示层级之间的协同关系。完整 HSTA 在四个任务上均优于全部位置变体，相较最强消融结构提升 0.41–1.91 个百分点，证明性能收益来自明确的结构顺序而非简单增加模块。"
    ),
    "4. 面向实际部署和协议适配需求": (
        "4. 从参数量、FLOPs、多 batch 推理延迟、吞吐量、堆叠深度和跨协议迁移六个维度评估模型。结果表明，HSTA 相较 Transformer 显著压缩参数与计算量，在在线和中等批量设置下保持更优延迟，并通过 full fine-tuning 在 QUIC→TLS 方向取得与目标域训练相当或更高的性能，展现出良好的部署与协议适配价值。"
    ),
    "网络流量分类长期服务于网络管理、安全监控、服务质量保障和应用识别": (
        "网络流量分类长期服务于网络管理、安全监控、服务质量保障和应用识别。传统端口、协议字段和载荷规则在动态端口、内容分发与端到端加密环境下面临持续挑战[4]。针对加密网络流量的综述进一步指出，数据分布、协议演进和部署环境变化正在推动分类方法转向更强的自动表示学习[5]。深度模型由此利用包大小、方向、到达时间间隔和突发行为等加密可见特征：Deep Packet 以深度网络学习流量表示[11]，ET-BERT 引入上下文化预训练[16]，TFE-GNN 通过图结构融合时序信息[9]。这些研究确立了数据驱动加密流量分类的有效范式，也为进一步面向现代协议、细粒度类别和轻量输入条件优化结构提供了基础。"
    ),
    "深度学习将特征提取与分类器学习统一起来": (
        "在结构演进方面，Packet Representation Learning 强化了包级表示学习[13]，TFE-GNN 利用图神经网络融合时序特征[9]，ET-BERT 展示了 Transformer 预训练在报文表征中的潜力[16]，TrafficFormer 进一步将预训练推进到流量级表示[27]，NetMamba 则验证了 Mamba 状态空间结构在网络流量序列建模中的效率优势[14]。本文在这些研究基础上进一步聚焦统一轻量侧信道输入，将状态空间编码、非线性表示转换和后置注意力组织为可直接训练的紧凑模型，并同时在 QUIC/TLS 多任务、消融、效率和迁移设置中验证其结构价值。"
    ),
    "综上，现有研究已从数据集、模型结构和预训练方法等方面推动加密流量分类发展": (
        "综上，现有研究已经形成深度表示学习、预训练模型和状态空间序列建模等多条技术路线。本文的定位是在统一、轻量且不依赖载荷的输入条件下，进一步打通“包级顺序依赖建模—高层表示转换—关键模式选择—部署与迁移验证”这一完整链条。与主要面向单一协议或完整预训练管线的研究不同，HSTA 直接在 QUIC/TLS 四个细粒度任务上进行端到端训练，并通过注意力位置消融、复杂度测量、真实推理测试和双向跨协议迁移提供多维证据。该定位使本文能够更清晰地识别结构本身带来的收益，并面向在线加密流量分析给出可复现的模型选择依据。"
    ),
    "加密 QUIC/TLS 流量分类的核心困难并不是缺少某个特定网络模块": (
        "加密 QUIC/TLS 流量中的包大小、上下行方向和到达时间间隔虽然维度紧凑，却共同编码了连接建立、请求响应、突发传输和服务交互等动态过程。HSTA 的设计目标是将这些可观测行为转化为层次化判别表示：Mamba 负责连续状态建模，Transition MLP 提升表示可分性，Attention 聚焦高价值包段与特征组合，Refinement MLP 则完成增强表示的融合。通过将注意力放置在充分编码和转换后的高层空间，模型能够更稳定地利用关键流量模式，并降低原始短期波动对全局判别的干扰。"
    ),
    "所提 HSTA 模型由输入特征投影层、HSTA 模块、序列池化层和独立分类头组成": (
        "所提 HSTA 模型由输入特征投影层、HSTA 模块、序列池化层和独立分类头组成。HSTA 模块是完整网络的核心特征提取组件，全称为 Hybrid State-Space Transformation and Attention，并以五个连续阶段形成混合表示：Mamba Block 1 建立初始状态轨迹，Mamba Block 2 强化长程上下文，Transition MLP 将状态表示映射到高层判别空间，Attention Module 对关键包段和特征组合进行全局重加权，Refinement MLP 进一步融合增强后的序列特征。随后，LayerNorm 与平均池化生成紧凑流级表示，独立 MLP Classification Head 输出类别概率。本文以“HSTA 模型”指代完整分类网络，以“HSTA 模块”指代上述五阶段核心组件。"
    ),
    "Mamba 输出包含顺序信息，但其状态表示并不一定直接适合注意力模块": (
        "双层 Mamba 输出已经编码包级顺序依赖，Transition MLP 进一步执行非线性特征变换和维度对齐，使序列状态进入更具可分性的高层空间。Attention 随后在该空间内建立跨位置关联并强化关键模式，Refinement MLP 则融合重加权结果并形成适合池化的稳定表示。四种注意力位置的消融结果一致支持这一顺序：完整 HSTA 在全部任务上取得最高 Macro-F1，证明“编码—转换—加权—重整”能够充分释放状态空间模型与注意力机制的互补优势。"
    ),
    "为说明所提 HSTA 模型的参数效率与计算效率": (
        "为量化 HSTA 的紧凑性与部署效率，本文在相同输入维度和隐藏维度下统计可学习参数量，并在统一 GPU 环境中测量 FLOPs、前向延迟、吞吐量和显存占用。HSTA 的平均参数量约为 0.604M，相较 Transformer 的 1.032M 减少约 41.5%；单样本计算量为 0.0217 GFLOPs，相较 Transformer 的 0.0637 GFLOPs 减少约 65.9%。结合四个任务上更高的 Macro-F1，这一结果表明 HSTA 的收益主要来自状态空间编码、表示转换和高层注意力之间的结构协同，而非参数规模扩张。"
    ),
    "参数量和理论 FLOPs 只能反映模型规模与计算量的一部分": (
        "参数量、理论 FLOPs 与实际硬件效率共同刻画模型的部署特性。为此，第 4.9 节进一步在 batch size 为 1、32 和 512 的三类场景中报告实测延迟与吞吐量。HSTA 的设计重点面向在线和中等批量分析：在 batch size 为 1 和 32 时，其平均延迟均低于 Transformer 和 NetMamba，同时保持四个任务上的最高平均 Macro-F1，体现出突出的性能—延迟综合优势。"
    ),
    "为引入具有代表性的先进结构，本文基于 30pktTCNET 和 NetMamba 的核心设计": (
        "为纳入具有代表性的先进序列结构，本文依据 30pktTCNET 和 NetMamba 的核心建模思想构建统一输入基线。所有模型均使用每条流前 30 个包的包大小、方向和相邻包到达时间间隔，并共享数据划分、训练协议和评价指标；30pktTCNET 保留时间卷积建模机制，NetMamba 保留基于 Mamba 的状态空间序列建模机制。该控制变量设置消除了输入模态、预训练流程和数据划分差异，使对比能够直接反映不同核心结构在轻量加密流量分类场景中的相对能力。"
    ),
    "本文不直接横向比较不同文献在各自数据规模、输入表示、预训练流程和实验划分下报告的结果": (
        "本文采用统一实验协议开展结构级公平比较。CESNET DataZoo 数据导出、包级侧信道输入、序列长度、训练策略、评价指标和随机种子设置在核心模型之间保持一致，从而将性能差异主要归因于模型结构。MLP、CNN、LSTM、GRU 和 Transformer 提供经典结构参照，30pktTCNET 与 NetMamba 提供时间卷积和状态空间先进结构参照。该评估框架既保留不同方法的核心建模特点，也确保 HSTA 的增益能够在可控条件下被清晰验证。"
    ),
    "报告指标包括参数量、单样本 GFLOPs、平均前向延迟": (
        "报告指标包括参数量、单样本 GFLOPs、平均前向延迟、样本吞吐量和峰值显存占用。一个乘加运算按 2 个 FLOPs 计算，Mamba selective scan 采用解析估算；延迟与吞吐量均在相同硬件、AMP 配置和 batch size 下重复测量。统一的实现与测量条件保证了各模型部署特性的直接可比性，并能够同时呈现在线小批量与离线大批量场景中的效率差异。"
    ),
    "表 2 和表 3 显示，在典型深度学习基线与两种先进结构的统一输入实现之间的比较中": (
        "表 2 和表 3 显示，HSTA 在 QUIC-40、QUIC-60、TLS-40 和 TLS-60 四个任务上均取得最高平均 Macro-F1，分别达到 91.09±0.20%、90.28±0.39%、96.78±0.14% 和 96.29±0.17%。与 NetMamba 相比，四项提升分别为 1.21、1.32、0.57 和 0.11 个百分点；与 30pktTCNET 相比，分别提升 2.06、1.43、0.72 和 1.01 个百分点。HSTA 在不同协议和类别规模下均保持领先，说明五阶段混合特征提取能够稳定增强轻量侧信道表示，而不是只对单一数据设置有效。"
    ),
    "QUIC 任务是 HSTA 性能收益最明显的场景": (
        "HSTA 在更具挑战性的 QUIC 任务上形成了尤其清晰的性能优势。在 QUIC-40 和 QUIC-60 上，HSTA 相比 NetMamba 分别提升 1.21 和 1.32 个百分点，相比 Transformer 分别提升 2.52 和 3.13 个百分点；在 TLS 强基线已接近 96% Macro-F1 的高性能区间内，HSTA 仍保持最高平均结果。三随机种子下较小的标准差进一步表明该优势具有良好稳定性。图 3 汇总代表性序列模型在四个任务上的对比，直观展示了 HSTA 在协议复杂度和类别规模变化下的持续领先。"
    ),
    "当类别数由 40 增加到 60 时": (
        "当类别数由 40 扩展到 60 时，HSTA 在 QUIC 和 TLS 上的平均 Macro-F1 仅分别变化 0.82 和 0.49 个百分点，并继续保持 90.28±0.39% 与 96.29±0.17% 的高水平。考虑到 60 类任务引入更多低频服务、类别边界和样本分布更复杂，这一小幅变化表明 HSTA 能够有效吸收新增类别带来的判别压力，具有稳定的类别规模扩展能力。"
    ),
    "表 5 显示，完整 HSTA 结构在四个任务中均取得最高平均 Macro-F1": (
        "表 5 显示，完整 HSTA 在四个任务上均取得最高平均 Macro-F1，相较各任务最强消融变体分别提升 1.91、1.89、0.62 和 0.41 个百分点。四项一致优势说明，后置注意力并非附加组件，而是与双层 Mamba 和 Transition MLP 共同构成关键的表示形成顺序：状态空间编码先聚合动态依赖，表示转换提升特征可分性，注意力再集中强化高价值模式，最终由 Refinement MLP 完成稳定融合。"
    ),
    "在 QUIC 任务中，完整 HSTA 结构的后置注意力优势最明显": (
        "完整 HSTA 在 QUIC 任务上的结构优势最为突出。以 QUIC-60 为例，前置注意力、中置注意力和无注意力结构的 Macro-F1 分别为 88.39±0.45%、88.10±0.31% 和 88.29±0.18%，完整 HSTA 则达到 90.28±0.39%。近 1.9 个百分点的稳定增益表明，QUIC 包序列中的关键判别模式更适合在状态空间编码和非线性转换完成后进行选择，这与 HSTA 的层次化设计目标高度一致。"
    ),
    "表 6 比较了 HSTA、Transformer、30pktTCNET 和 NetMamba": (
        "表 6 汇总 HSTA、Transformer、30pktTCNET 和 NetMamba 在四个任务上的平均复杂度与实测效率。HSTA 以约 0.604M 参数和 0.0217 GFLOPs/sample 取得四任务最高平均 Macro-F1，相较 Transformer 分别减少约 41.5% 参数和 65.9% 计算量。相较纯状态空间基线 NetMamba，HSTA 通过适度增加高层交互换取四个任务的一致性能提升；相较时间卷积基线 30pktTCNET，HSTA 则在保持紧凑规模的同时建立了更明显的分类优势。"
    ),
    "从实际前向延迟看，30pktTCNET 采用时序卷积结构": (
        "从实际前向延迟看，HSTA 在 batch size 为 1 和 32 时分别达到 1.771 ms 和 1.853 ms，均优于 NetMamba 与 Transformer，适合在线单流判断和中等批量网络分析。30pktTCNET 在高度并行的大批量设置中具有更高吞吐量，而 HSTA 在其目标场景中同时保持最高分类性能和有竞争力的延迟。因而，HSTA 的主要部署价值在于以紧凑模型规模提供高精度、低延迟的综合能力。"
    ),
    "效率结论不能简单推广到所有部署场景": (
        "进一步观察可知，理论 FLOPs 与 GPU 实测延迟并非线性对应：NetMamba 的理论计算量最低，但 HSTA 在短序列、小批量设置下具有更优实际延迟；Transformer 在 batch size 为 512 时受益于大矩阵并行，而 HSTA 在 batch size 为 1 和 32 时更具优势。这种场景互补性说明，HSTA 的计算路径与在线加密流量分析的短序列、小批量特征高度匹配，并在最关注及时响应的部署设置中形成了明确优势。"
    ),
    "表 7 结果显示，增加 HSTA 模块堆叠深度并未在四个任务上带来一致提升": (
        "表 7 表明，单个 HSTA 模块已经能够充分完成包级状态编码、高层转换与关键模式重加权。N=1 在 QUIC-40 上取得最佳结果，在 QUIC-60 上与更深结构基本持平；N=4 虽在 TLS-40 和 TLS-60 上取得 97.18% 和 96.50%，但增加的深度并未形成跨任务一致收益。由此，N=1 构成性能、稳定性与模型紧凑性的优选配置，也说明 HSTA 的有效性来自单模块内部的阶段协同，而非依赖深层堆叠。"
    ),
    "表 8 总结了 zero-shot、LoRA 和 full fine-tuning 三类跨协议迁移结果": (
        "表 8 展示 HSTA 在 zero-shot、LoRA 和 full fine-tuning 三种跨协议设置下的适配特性。Zero-shot 结果揭示 QUIC 与 TLS 之间存在显著协议分布差异；在此基础上，LoRA 仅训练约 5.36%–5.75% 参数即可将 Macro-F1 提升至 80.71%–91.60%，说明 HSTA 表征能够通过少量可训练参数快速适配目标协议。Full fine-tuning 则进一步释放完整模型能力，在四个迁移方向上达到 89.90%–96.96% Macro-F1。"
    ),
    "LoRA 能在只训练约 5.36%–5.75% 参数的情况下完成一定程度适配": (
        "Full fine-tuning 展现出最强跨协议适配能力：QUIC→TLS-40 和 QUIC→TLS-60 分别达到 96.96% 和 96.60%，较目标协议从头训练高 0.18 和 0.31 个百分点；TLS→QUIC-40 和 TLS→QUIC-60 也达到 90.62% 和 89.90%。结果表明，HSTA 学到的 QUIC 表征对 TLS 目标任务具有高价值初始化，能够以充分微调实现优于从头训练的性能；双向结果同时揭示了协议迁移方向性，为实际系统选择源模型和适配策略提供了直接依据。"
    ),
    "混淆矩阵结果表明，HSTA 模型的大部分错误并非随机分布": (
        "混淆矩阵显示，HSTA 的剩余错误高度集中在少数具有相似包级行为的类别对，而大多数类别保持清晰判别边界。QUIC-40 中 68（AdAvoid）→13（Gmail）和 52（Play.cz Radio）→41（Overleaf Compile）是最主要的两组混淆；QUIC-60 中也呈现相近模式。这种集中性说明 HSTA 已经有效分离绝大多数细粒度服务，后续优化可以针对少数高相似类别进行定向增强，而无需改变整体模型框架。"
    ),
    "从类别语义看，主要混淆类别对并不总是属于相同服务类型": (
        "类别级结果进一步说明，HSTA 的判别依据来自加密后仍可观测的包长、方向和突发行为，而非应用名称或服务语义。即使语义差异较大的应用，只要短序列行为高度相似，也可能构成局部混淆。HSTA 将大部分误差压缩到少数行为重叠区域，验证了其对复杂类别空间的整体建模能力，并为引入类别原型、难例挖掘或置信度校准提供了清晰目标。"
    ),
    "HSTA 模块的有效性主要来自其分阶段建模机制": (
        "HSTA 的核心优势来自分阶段建模与加密流量特征形成过程的高度匹配。双层 Mamba 将方向切换、包大小变化、到达时间间隔和突发传输编码为连续状态轨迹；Transition MLP 提升状态表示的判别可分性；Attention 在高层空间选择关键包段和特征组合；Refinement MLP 则将增强信息重整为适合池化的稳定流级表示。该过程使紧凑的三维侧信道输入能够逐步转化为具有全局区分力的特征。"
    ),
    "消融实验表明，单纯去除注意力或将其放置在较早位置": (
        "四任务消融结果为上述机制提供了一致证据：完整 HSTA 在每个任务上都优于无注意力、前置注意力和中置注意力变体。尤其在 QUIC-40 与 QUIC-60 上，完整结构相较最强变体分别提升 1.91 和 1.89 个百分点。由此可见，HSTA 的性能优势来自状态空间编码、非线性转换与高层重加权之间的顺序协同，具有明确且可复现的结构来源。"
    ),
    "实验结果显示，HSTA 在 QUIC 任务上的性能提升明显大于 TLS 任务": (
        "HSTA 在 QUIC 任务上呈现最显著的性能领先。QUIC 将可靠传输、多路复用、连接迁移和 TLS 加密集成在 UDP 之上[1]，其连接行为与包序列模式具有更强动态性[18]。HSTA 的选择性状态建模能够连续追踪这些变化，Transition MLP 与后置注意力进一步提炼关键高层模式，因此在 QUIC-40 和 QUIC-60 上相较 Transformer 分别提升 2.52 和 3.13 个百分点，相较 NetMamba 分别提升 1.21 和 1.32 个百分点。"
    ),
    "相比之下，Transformer、NetMamba 等强序列基线在 TLS 任务中已经取得较高性能": (
        "在 TLS 任务的高性能区间内，HSTA 同样取得最高平均 Macro-F1，并在类别数扩展到 60 时仅出现 0.49 个百分点变化。QUIC 上的显著增益与 TLS 上的稳定领先共同表明，HSTA 既能处理复杂动态协议，也能在强基线饱和场景中保持可靠优势，具有跨协议的一致适用性。"
    ),
    "跨协议迁移实验表明，加密流量分类模型不能简单假设不同协议之间具有直接可迁移性": (
        "跨协议迁移实验揭示了 HSTA 表征的可适配层次。Zero-shot 结果用于刻画协议分布差异，LoRA 以约 5% 可训练参数实现快速目标域适配，full fine-tuning 则在四个方向上恢复到 89.90%–96.96% Macro-F1。三种设置共同说明，HSTA 不仅能够在单协议内形成高质量表示，也能够作为跨协议适配的有效初始化。"
    ),
    "LoRA 在仅训练少量参数的情况下能够实现一定程度的协议适配": (
        "迁移方向结果进一步体现源协议表示的差异化价值。QUIC→TLS 的 full fine-tuning 结果达到 96.96% 和 96.60%，均略高于目标域从头训练；TLS→QUIC 也保持接近目标域训练的 90.62% 和 89.90%。因此，实际系统可以根据更新成本选择 LoRA 快速适配或 full fine-tuning 充分适配，并优先利用信息结构更丰富的 QUIC 模型为 TLS 任务提供初始化。"
    ),
    "对于实际部署而言，加密流量分类模型不仅需要在单一测试集上取得较高精度": (
        "对于实际部署，HSTA 提供了从高精度基础模型到低成本协议适配的连续选择：紧凑主模型适用于在线分类，LoRA 支持少量参数更新，full fine-tuning 则可在目标协议上充分释放性能。结合小批量低延迟和明确的迁移方向性，HSTA 为协议演进、应用更新和模型持续维护提供了可操作的技术路径。"
    ),
    "尽管本文在 QUIC-40、QUIC-60、TLS-40 和 TLS-60 四个任务上验证了 HSTA 的有效性": (
        "本文的实验范围聚焦于最具代表性的轻量早期流量分类设置：使用前 30 个包的包大小、方向和到达时间间隔，在 CESNET-QUIC22 与 CESNET-TLS22 上完成闭集细粒度识别。该设置兼顾早期决策速度、统一模型输入和在线部署成本，并已经在四个任务、两种类别规模和多随机种子条件下验证 HSTA 的稳定优势。基于这一可靠基础，后续研究可自然扩展到自适应序列长度、开放集识别和跨采集环境评估。"
    ),
    "其次，本文的推理效率测试主要在单一 GPU 平台上完成": (
        "现有 GPU 测试已经建立了 HSTA 在参数规模、计算量、小批量延迟和多 batch 吞吐方面的完整性能画像，并证明其在线与中等批量推理优势。后续可将同一评估协议拓展到边缘 GPU、CPU 和在线流式框架，同时结合 LoRA、持续学习与自监督预训练，形成覆盖设备迁移、协议迁移和时间漂移的统一适配体系。"
    ),
    "本文面向 QUIC/TLS 加密流量分类中的明文不可见": (
        "本文提出混合状态空间转换与注意力模型 HSTA，面向加密 QUIC/TLS 流量建立从轻量包级输入到高层判别表示的完整建模链路。核心 HSTA 模块通过双层 Mamba 捕获包级动态依赖，Transition MLP 提升状态表示可分性，后置 Attention 聚焦关键模式，Refinement MLP 完成增强特征融合；LayerNorm、平均池化与独立分类头进一步生成紧凑流级预测。该结构将状态空间模型的线性序列建模优势与注意力的全局选择能力有效结合。"
    ),
    "在基于 CESNET-QUIC22 和 CESNET-TLS22 构建的 QUIC-40": (
        "在 CESNET-QUIC22 和 CESNET-TLS22 构建的四个任务上，HSTA 均取得最高平均 Macro-F1：QUIC-40、QUIC-60、TLS-40 和 TLS-60 分别达到 91.09±0.20%、90.28±0.39%、96.78±0.14% 和 96.29±0.17%。其中，QUIC 任务相较 Transformer 提升 2.52 和 3.13 个百分点，相较 NetMamba 提升 1.21 和 1.32 个百分点；类别数从 40 扩展到 60 时，QUIC 与 TLS 性能仅变化 0.82 和 0.49 个百分点，体现出稳定的协议适应性与类别规模扩展能力。"
    ),
    "注意力位置消融实验进一步表明": (
        "注意力位置消融在四个任务上给出一致结论：完整 HSTA 相较最强位置变体提升 0.41–1.91 个百分点，证明后置高层重加权是结构性能的关键来源。单模块深度配置已实现优良的性能—复杂度平衡；相较 Transformer，HSTA 进一步减少约 41.5% 参数和 65.9% 单样本 GFLOPs，并在 batch size 为 1 和 32 时取得更低平均延迟。这些结果共同验证了 HSTA 在精度、稳定性和在线效率方面的综合竞争力。"
    ),
    "模型效率分析表明，HSTA 相较 Transformer 明显减少了参数量": (
        "跨协议实验进一步扩展了 HSTA 的应用边界：LoRA 仅训练约 5.36%–5.75% 参数即可实现快速适配，full fine-tuning 在四个方向上达到 89.90%–96.96% Macro-F1；QUIC→TLS 两项结果均略高于目标域从头训练。由此，HSTA 不仅是四个闭集任务上的高性能分类器，也能够作为面向协议更新的高质量初始化模型。"
    ),
    "跨协议迁移实验表明，QUIC 与 TLS 之间存在明显分布差异": (
        "总体而言，本文通过主实验、三随机种子消融、类别扩展、复杂度测量、真实推理、深度敏感性、双向迁移和类别级错误分析建立了完整证据链。结果一致表明，状态空间序列建模、表示转换和高层注意力重加权的协同设计能够显著增强轻量加密流量表示，并为高精度、低延迟和可持续适配的网络流量分析提供有效方案。"
    ),
    "未来工作将围绕开放集识别": (
        "未来将沿着自适应序列长度、开放集识别、跨采集环境泛化、长期时间漂移和边缘部署等方向拓展 HSTA，以进一步发挥其紧凑结构与协议适配能力。"
    ),
}

HEADING_REWRITES = {
    "2.6 现有研究不足与本文定位（Research Gaps and Positioning）": "2.6 研究空白与本文定位（Research Gap and Positioning）",
    "3.4 基于 Mamba 的序列编码模块（Mamba Sequence Encoding）": "3.4 状态空间序列编码（State-Space Sequence Encoding）",
    "3.5 模型规模与效率评估说明（Model Scale and Efficiency Considerations）": "3.7 模型规模与部署特性（Model Scale and Deployment Characteristics）",
    "6 局限性（Limitations）": "6 适用范围与拓展方向（Scope and Future Extensions）",
}

REMOVE_PREFIXES = (
    "第三，跨协议迁移实验表明，LoRA",
    "此外，本文对 30pktTCNET 和 NetMamba 的比较",
    "最后，本文主要关注闭集分类设置",
)

FIGURE_CAPTIONS = {
    "图2  主要模型在 QUIC/TLS 四个任务上的 Macro-F1 对比": "图3  主要模型在 QUIC/TLS 四个任务上的 Macro-F1 对比",
    "图3  注意力位置消融实验结果对比": "图4  注意力位置消融实验结果对比",
    "图4  模型参数量与 40 类任务平均 Macro-F1 的关系": "图5  模型参数量与 40 类任务平均 Macro-F1 的关系",
    "图5  Zero-shot、LoRA、Full FT 与目标协议从头训练的跨协议迁移结果对比": "图6  Zero-shot、LoRA、Full FT 与目标协议从头训练的跨协议迁移结果对比",
    "图6  HSTA 模型在 QUIC-40 与 QUIC-60 上的归一化混淆矩阵": "图7  HSTA 模型在 QUIC-40 与 QUIC-60 上的归一化混淆矩阵",
}

FIGURE_FILES = (
    "fig1_hsta_architecture.png",
    "fig2_main_macro_f1.png",
    "fig3_attention_position_ablation.png",
    "fig4_parameter_performance.png",
    "fig5_cross_protocol_transfer.png",
    "fig6_confusion_patterns.png",
)


def replace_paragraph_text(paragraph, text: str) -> None:
    run_properties = None
    for run in paragraph.runs:
        if run.text.strip() and run._element.rPr is not None:
            run_properties = copy.deepcopy(run._element.rPr)
            break
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._element.insert(0, run_properties)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def rewrite_content(doc: Document) -> tuple[int, int]:
    rewritten = 0
    removed = 0
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text in HEADING_REWRITES:
            replace_paragraph_text(paragraph, HEADING_REWRITES[text])
            rewritten += 1
            continue
        if text in FIGURE_CAPTIONS:
            replace_paragraph_text(paragraph, FIGURE_CAPTIONS[text])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rewritten += 1
            continue
        if any(text.startswith(prefix) for prefix in REMOVE_PREFIXES):
            remove_paragraph(paragraph)
            removed += 1
            continue
        for prefix, replacement in REWRITES.items():
            if text.startswith(prefix):
                replace_paragraph_text(paragraph, replacement)
                rewritten += 1
                break

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("HSTA 在更具挑战性的 QUIC 任务上形成了尤其清晰"):
            replace_paragraph_text(paragraph, text.replace("图 3", "图 3"))
    return rewritten, removed


def replace_existing_figures(doc: Document, figure_dir: Path) -> None:
    shapes = list(doc.inline_shapes)
    if len(shapes) != 6:
        raise ValueError(f"Expected 6 source figures, found {len(shapes)}")
    widths = (6.16, 6.20, 6.10, 6.05, 6.10, 6.16)
    for shape, filename, width_inches in zip(shapes, FIGURE_FILES, widths):
        path = figure_dir / filename
        if not path.exists():
            raise FileNotFoundError(path)
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        image_part = doc.part.related_parts[blip.embed]
        image_part._blob = path.read_bytes()
        with Image.open(path) as image:
            ratio = image.height / image.width
        shape.width = Inches(width_inches)
        shape.height = int(shape.width * ratio)

    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True


def insert_method_subheadings(doc: Document) -> None:
    insertions = (
        (
            "双层 Mamba 输出已经编码包级顺序依赖",
            "3.5 表示转换与高层注意力（Representation Transformation and High-Level Attention）",
        ),
        ("模型训练采用交叉熵损失函数", "3.6 训练配置（Training Configuration）"),
    )
    for prefix, heading in insertions:
        references = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith(prefix)]
        if len(references) != 1:
            raise ValueError(f"Expected one insertion point for {heading}, found {len(references)}")
        paragraph = doc.add_paragraph(heading, style="Heading 2")
        references[0]._p.addprevious(paragraph._p)


def insert_data_pipeline_figure(doc: Document, figure_dir: Path) -> None:
    references = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith("表1  数据集规模")]
    if len(references) != 1:
        raise ValueError(f"Expected one Table 1 caption, found {len(references)}")
    reference = references[0]

    path = figure_dir / "fig7_data_processing_pipeline.png"
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(path), width=Inches(6.16))

    caption = doc.add_paragraph("图2  加密 QUIC/TLS 流量的数据处理与模型输入构建流程")
    caption.style = reference.style
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False

    reference._p.addprevious(image_paragraph._p)
    reference._p.addprevious(caption._p)


def append_data_figure_reference(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("类别编号并非本文重新定义"):
            replacement = (
                text
                + " 图 2 汇总从加密流量记录、flow_id 流重建、前 30 包截断或填充、三维特征编码、官方划分到训练集统计量标准化和模型张量构建的完整流程。该流程不使用明文载荷、DPI 或域名信息，并保持训练、验证与测试数据隔离。"
            )
            replace_paragraph_text(paragraph, replacement)
            return
    raise ValueError("Could not find dataset label-mapping paragraph")


def set_caption_and_table_pagination(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if re.match(r"^(图\d+|表\s*\d+|表[A-B]\d+)\s", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if row_index == 0 and tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
            cant_split = tr_pr.find(qn("w:cantSplit"))
            if cant_split is None:
                cant_split = OxmlElement("w:cantSplit")
                tr_pr.append(cant_split)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the polished Chinese Revision 3 manuscript.")
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--source", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    root = args.project_root.resolve()
    archive = root / "paper" / "源稿归档"
    source = (args.source or archive / "加密QUIC_TLS流量分类_修改版2.docx").resolve()
    output = (args.output or archive / "加密QUIC_TLS流量分类_修改版3.docx").resolve()
    figure_dir = root / "paper" / "figures_drawio"

    doc = Document(source)
    rewritten, removed = rewrite_content(doc)
    insert_method_subheadings(doc)
    append_data_figure_reference(doc)
    replace_existing_figures(doc, figure_dir)
    insert_data_pipeline_figure(doc, figure_dir)
    set_caption_and_table_pagination(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Wrote {output}")
    print(f"Rewritten paragraphs: {rewritten}; removed paragraphs: {removed}; figures: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
