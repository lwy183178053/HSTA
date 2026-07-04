from __future__ import annotations

import importlib.util
import io
import re
import sys
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "paper" / "submission_jnca"
BILINGUAL = ROOT / "paper" / "jnca_bilingual_latex"
ENGLISH_OVERLEAF_ZIP = ROOT / "paper" / "jnca_english_overleaf_package.zip"
UPLOAD_BUNDLE = SUBMISSION / "JNCA_recommended_upload_bundle.zip"
OVERLEAF_DIRECT_ZIP = SUBMISSION / "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip"
EDITORIAL_SOURCE_ZIP = SUBMISSION / "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip"
FIGURE_UPLOAD_ZIP = SUBMISSION / "FIGURES_FOR_UPLOAD.zip"
GRAPHICAL_ABSTRACT = SUBMISSION / "Graphical_Abstract.png"
GRAPHICAL_ABSTRACT_TIF = SUBMISSION / "Graphical_Abstract.tif"
FIGURE_CAPTIONS_DOCX = SUBMISSION / "Figure_Captions.docx"
OVERLEAF_IMPORT_GUIDE = SUBMISSION / "OVERLEAF_IMPORT_GUIDE.md"
GFA_COMPLIANCE_CHECK = SUBMISSION / "JNCA_GFA_COMPLIANCE_CHECK.md"
PLACEHOLDER_REPLACEMENT_CHECK = SUBMISSION / "FINAL_PLACEHOLDER_REPLACEMENT_CHECK.md"
IMPORT_FORMAT_CHECK = SUBMISSION / "SUBMISSION_IMPORT_FORMAT_CHECK.md"
LATEX_TROUBLESHOOTING = SUBMISSION / "LATEX_IMPORT_TROUBLESHOOTING.md"
SUBMISSION_SCRIPT = ROOT / "scripts" / "generate_jnca_submission.py"
OFFICIAL_ELSARTICLE_TEMPLATE = SUBMISSION / "official_template" / "elsarticle" / "elsarticle"
OFFICIAL_ELSARTICLE_BST = OFFICIAL_ELSARTICLE_TEMPLATE / "elsarticle-num.bst"
MANUSCRIPT_TITLE = "Hybrid State-Space Transition-Attention Modeling for Lightweight Encrypted QUIC/TLS Traffic Classification"
MANUSCRIPT_KEYWORDS = "Encrypted traffic classification; QUIC; TLS; Mamba; state-space model; cross-protocol transfer"
CORE_PROPERTY_AUTHOR = "Author metadata placeholder"
PDF_FORBIDDEN_RAW_MARKERS = [
    b"/Author",
    b"/Subject",
    b"/Keywords",
    b"Author metadata placeholder",
    b"python-docx",
]


def fail(message: str, failures: list[str]) -> None:
    failures.append(message)
    print(f"[FAIL] {message}")


def ok(message: str) -> None:
    print(f"[OK] {message}")


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def source_text_hygiene_issues(data: bytes, *, ascii_only: bool) -> list[str]:
    issues: list[str] = []
    if data.startswith(b"\xef\xbb\xbf"):
        issues.append("UTF-8 BOM")
    if b"\x00" in data:
        issues.append("NUL byte")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        issues.append(f"invalid UTF-8 at byte {exc.start}")
        return issues
    if "\r\n" in text and re.search(r"(?<!\r)\n|\r(?!\n)", text):
        issues.append("mixed newline style")
    if ascii_only:
        non_ascii = sorted({ch for ch in text if ord(ch) > 127})
        if non_ascii:
            issues.append(f"non-ASCII characters: {non_ascii[:12]}")
    return issues


def check_source_text_hygiene(path: Path, failures: list[str], *, ascii_only: bool) -> None:
    if not check_exists(path, failures):
        return
    issues = source_text_hygiene_issues(path.read_bytes(), ascii_only=ascii_only)
    if issues:
        fail(f"{path.name} has source text hygiene issues: {issues}", failures)
    else:
        scope = "ASCII/UTF-8" if ascii_only else "UTF-8"
        ok(f"{path.name} source text hygiene is {scope}-safe")


def normalize_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def load_submission_module():
    spec = importlib.util.spec_from_file_location("generate_jnca_submission_for_validation", SUBMISSION_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load submission generator: {SUBMISSION_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def table_rows(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def docx_zip_text(path: Path, pattern: str) -> str:
    with zipfile.ZipFile(path) as zf:
        return "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist()
            if re.fullmatch(pattern, name)
        )


def extract_section_text(doc: Document, heading: str) -> list[str]:
    capture = False
    items: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            if text == heading:
                capture = True
                continue
            if capture:
                break
        elif capture:
            items.append(text)
    return items


def tex_block(text: str, env: str) -> str:
    match = re.search(rf"\\begin\{{{re.escape(env)}\}}(.*?)\\end\{{{re.escape(env)}\}}", text, re.S)
    return match.group(1).strip() if match else ""


def tex_title(text: str) -> str:
    match = re.search(r"\\title\{([^}]*)\}", text)
    return normalize_spaces(match.group(1)) if match else ""


def tex_plain_text(text: str) -> str:
    text = re.sub(r"~\\cite\{[^}]+\}", "", text)
    text = re.sub(r"\\cite\{[^}]+\}", "", text)
    text = text.replace(r"\%", "%")
    text = text.replace(r"\pm", "±")
    text = re.sub(r"\$([^$]*)\$", r"\1", text)
    text = text.replace(r"\sep", ";")
    text = re.sub(r"\\item\s+", "", text)
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^]]*\])?(?:\{[^}]*\})?", "", text)
    text = text.replace("{", "").replace("}", "")
    return normalize_spaces(text)


def tex_keywords(text: str) -> list[str]:
    return [normalize_spaces(item) for item in tex_block(text, "keyword").split(r"\sep") if normalize_spaces(item)]


def tex_highlights(text: str) -> list[str]:
    return [normalize_spaces(item) for item in re.findall(r"\\item\s+(.+)", tex_block(text, "highlights"))]


def image_size(path: Path) -> tuple[int, int] | None:
    try:
        from PIL import Image
    except ModuleNotFoundError:
        return None
    with Image.open(path) as img:
        return img.size


def check_image_min_size(path: Path, failures: list[str], *, min_width: int, min_height: int, label: str) -> None:
    size = image_size(path)
    if size is None:
        fail(f"{label} size not inspected because Pillow is unavailable", failures)
        return
    if size[0] >= min_width and size[1] >= min_height:
        ok(f"{label} size is {size[0]}x{size[1]} px")
    else:
        fail(f"{label} size is {size[0]}x{size[1]} px, expected at least {min_width}x{min_height}", failures)


def zip_entry_hygiene_issues(names: list[str]) -> list[str]:
    issues: list[str] = []
    seen: set[str] = set()
    system_names = {".DS_Store", "Thumbs.db", "desktop.ini"}
    for name in names:
        if name in seen:
            issues.append(f"duplicate entry: {name}")
        seen.add(name)
        normalized = name.replace("\\", "/")
        parts = [part for part in normalized.split("/") if part]
        base = parts[-1] if parts else ""
        lower = normalized.lower()
        if "\\" in name:
            issues.append(f"backslash path separator: {name}")
        if name.startswith(("/", "\\")) or re.match(r"^[A-Za-z]:", name):
            issues.append(f"absolute path: {name}")
        if ".." in parts:
            issues.append(f"path traversal: {name}")
        if lower.startswith("__macosx/") or base in system_names:
            issues.append(f"system/hidden file: {name}")
        try:
            name.encode("ascii")
        except UnicodeEncodeError:
            issues.append(f"non-ASCII zip path: {name}")
    return issues


def check_zip_entry_hygiene(path: Path, zf: zipfile.ZipFile, failures: list[str]) -> None:
    issues = zip_entry_hygiene_issues(zf.namelist())
    if issues:
        fail(f"{path.name} has unsafe or nonportable zip entries: {issues[:8]}", failures)
    else:
        ok(f"{path.name} has clean zip entry paths")


def read_zip_text(zf: zipfile.ZipFile, name: str) -> str:
    return zf.read(name).decode("utf-8")


def check_zip_readme_consistency(path: Path, zf: zipfile.ZipFile, names: set[str], failures: list[str]) -> None:
    if "README.md" not in names:
        fail(f"{path.name} has no README.md for import guidance", failures)
        return
    readme = read_zip_text(zf, "README.md")
    issues: list[str] = []
    if "main.tex" not in readme:
        issues.append("README does not name main.tex")
    if "pdfLaTeX" not in readme and path.name != "jnca_chinese_xelatex_package.zip":
        issues.append("README does not name pdfLaTeX")
    if path.name == "jnca_chinese_xelatex_package.zip" and "XeLaTeX" not in readme:
        issues.append("README does not name XeLaTeX")
    if "elsarticle.cls" in readme and "expected from" not in readme:
        issues.append("README mentions elsarticle.cls without platform-dependency wording")
    if "FIGURES_FOR_UPLOAD.zip" in readme and "not for compiling" not in readme and "separate artwork" not in readme:
        issues.append("README mentions FIGURES_FOR_UPLOAD.zip without artwork-only wording")

    mentioned_files = sorted(set(re.findall(r"`([^`]+\.[A-Za-z0-9]+)`", readme)))
    missing_mentions = [
        item
        for item in mentioned_files
        if "/" not in item
        and "\\" not in item
        and item not in names
        and item
        not in {
            "elsarticle.cls",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_recommended_upload_bundle.zip",
            "FIGURES_FOR_UPLOAD.zip",
        }
        and not item.startswith("paper/")
        and not item.startswith("../")
    ]
    if missing_mentions:
        issues.append(f"README mentions files not present in the zip: {missing_mentions}")

    if path.name in {"jnca_english_overleaf_package.zip", "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip", "overleaf_jnca_package.zip"}:
        if "figures/" not in readme:
            issues.append("README does not explain figures/ layout")
        if "JNCA_recommended_upload_bundle.zip" not in readme:
            issues.append("README does not warn against importing the handoff bundle")
    if path.name == "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip":
        if "flat" not in readme.lower() or "figures/" in readme:
            issues.append("README does not clearly describe the flat Editorial Manager layout")
        png_entries = sorted(name for name in names if name.endswith(".png"))
        if not png_entries:
            issues.append("flat source README exists but no root-level PNG figures are present")
    if path.name == "jnca_bilingual_latex_package.zip":
        if "main_cn.tex" not in readme or "XeLaTeX" not in readme:
            issues.append("bilingual README does not explain the Chinese XeLaTeX entry point")

    if issues:
        fail(f"{path.name} README/package consistency issues: {issues}", failures)
    else:
        ok(f"{path.name} README is consistent with the package layout")


def check_exists(path: Path, failures: list[str]) -> bool:
    if path.exists() and path.stat().st_size > 0:
        ok(f"{path.relative_to(ROOT)} exists")
        return True
    fail(f"{path.relative_to(ROOT)} is missing or empty", failures)
    return False


def check_pdf(failures: list[str]) -> None:
    pdf = SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_EN.pdf"
    docx = SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_EN.docx"
    if check_exists(pdf, failures) and pdf.read_bytes()[:5] == b"%PDF-":
        ok("English PDF header is %PDF-")
    else:
        fail("English PDF header check failed", failures)
        return
    pdf_bytes = pdf.read_bytes()
    if b"%%EOF" in pdf_bytes[-2048:]:
        ok("English PDF has EOF marker")
    else:
        fail("English PDF missing EOF marker near file end", failures)
    leaked_markers = [marker.decode("latin1") for marker in PDF_FORBIDDEN_RAW_MARKERS if marker in pdf_bytes]
    if leaked_markers:
        fail(f"English PDF has obvious raw metadata/tool markers: {leaked_markers}", failures)
    else:
        ok("English PDF has no obvious raw author/tool metadata markers")
    page_count = len(re.findall(rb"/Type\s*/Page\b", pdf_bytes))
    if page_count >= 5:
        ok(f"English PDF has {page_count} page objects")
    else:
        fail(f"English PDF page object count is {page_count}, expected at least 5", failures)
    stream_count = len(re.findall(rb"stream\r?\n", pdf_bytes))
    if stream_count >= 10:
        ok(f"English PDF has {stream_count} content streams")
    else:
        fail(f"English PDF content stream count is {stream_count}, expected at least 10", failures)
    font_count = len(re.findall(rb"/Font\b", pdf_bytes))
    if font_count >= 1:
        ok(f"English PDF has {font_count} font resource markers")
    else:
        fail("English PDF has no font resource marker", failures)
    image_count = len(re.findall(rb"/Image\b", pdf_bytes))
    if image_count >= 4:
        ok(f"English PDF has {image_count} image object markers")
    else:
        fail(f"English PDF image object count is {image_count}, expected at least 4", failures)
    decoded_streams = 0
    decoded_bytes = 0
    for stream in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", pdf_bytes, re.S):
        try:
            decoded = zlib.decompress(stream)
        except zlib.error:
            continue
        decoded_streams += 1
        decoded_bytes += len(decoded)
    if decoded_streams >= 10 and decoded_bytes > 100_000:
        ok(f"English PDF has {decoded_streams} decodable streams ({decoded_bytes} bytes)")
    else:
        fail(f"English PDF decodable stream evidence is weak: {decoded_streams} streams, {decoded_bytes} bytes", failures)
    if docx.exists():
        if pdf.stat().st_mtime >= docx.stat().st_mtime:
            ok("English PDF is not older than the English Word manuscript")
        else:
            fail("English PDF is older than the English Word manuscript; regenerate PDF", failures)


def check_english_word_review_format(path: Path, doc: Document, failures: list[str]) -> None:
    normal = doc.styles["Normal"]
    size = normal.font.size.pt if normal.font.size is not None else None
    if size == 12:
        ok(f"{path.name} Normal style is 12 pt")
    else:
        fail(f"{path.name} Normal style size is {size}, expected 12 pt", failures)
    if normal.font.name == "Times New Roman":
        ok(f"{path.name} Normal style uses Times New Roman")
    else:
        fail(f"{path.name} Normal style font is {normal.font.name!r}, expected Times New Roman", failures)
    if normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE:
        ok(f"{path.name} Normal style uses double line spacing")
    else:
        fail(f"{path.name} Normal style does not use double line spacing", failures)

    body_paragraphs = [p for p in doc.paragraphs if p.text.strip() and not p.style.name.startswith("Heading")][:20]
    spacing_ok = all(p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE for p in body_paragraphs)
    if body_paragraphs and spacing_ok:
        ok(f"{path.name} body paragraphs use double line spacing")
    else:
        fail(f"{path.name} body paragraphs are not consistently double spaced", failures)

    document_xml = docx_zip_text(path, r"word/document\.xml")
    footer_xml = docx_zip_text(path, r"word/footer\d+\.xml")
    if 'w:lnNumType' in document_xml and 'w:restart="continuous"' in document_xml:
        ok(f"{path.name} has continuous line numbering XML")
    else:
        fail(f"{path.name} missing continuous line numbering XML", failures)
    if "PAGE" in footer_xml and "fldCharType" in footer_xml:
        ok(f"{path.name} has page-number field in footer")
    else:
        fail(f"{path.name} missing page-number field in footer", failures)

    abstract_text = " ".join(extract_section_text(doc, "Abstract"))
    abstract_words = len(re.findall(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)?", abstract_text))
    if 0 < abstract_words <= 250:
        ok(f"{path.name} abstract word count is {abstract_words} (<=250)")
    else:
        fail(f"{path.name} abstract word count is {abstract_words}, expected 1-250", failures)
    keywords_text = " ".join(extract_section_text(doc, "Keywords"))
    keywords = [item.strip() for item in keywords_text.split(";") if item.strip()]
    if 1 <= len(keywords) <= 6:
        ok(f"{path.name} keyword count is {len(keywords)}")
    else:
        fail(f"{path.name} keyword count is {len(keywords)}, expected 1-6", failures)


def check_english_word_structure_hygiene(path: Path, doc: Document, failures: list[str]) -> None:
    title_paragraphs = [normalize_spaces(p.text) for p in doc.paragraphs if p.style.name == "Title" and p.text.strip()]
    if title_paragraphs == [MANUSCRIPT_TITLE]:
        ok(f"{path.name} has one controlled title paragraph")
    else:
        fail(f"{path.name} title paragraphs are {title_paragraphs}, expected {[MANUSCRIPT_TITLE]}", failures)

    expected_headings = [
        "Abstract",
        "Keywords",
        "Highlights",
        "1 Introduction",
        "2 Background and Related Work",
        "3 Problem Definition and Motivation",
        "4 Proposed Method",
        "5 Experimental Setup",
        "6 Results and Analysis",
        "7 Discussion",
        "8 Threats to Validity",
        "9 Conclusion",
        "Declarations",
        "References",
        "Appendix A. QUIC label-to-service mapping",
        "Appendix B. TLS label-to-service mapping",
    ]
    headings = [normalize_spaces(p.text) for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip()]
    if headings == expected_headings:
        ok(f"{path.name} has the expected heading sequence")
    else:
        fail(f"{path.name} heading sequence is {headings}, expected {expected_headings}", failures)

    expected_tables = [
        "Table 1 Dataset and task statistics",
        "Table 2 Adapted advanced baselines under the unified input setting",
        "Table 3 Main results on 40-class tasks and adapted advanced baselines (%)",
        "Table 4 Main results on 60-class tasks and adapted advanced baselines (%)",
        "Table 5 Performance change from 40 to 60 classes (Macro-F1, %)",
        "Table 6 Attention-position ablation results (Macro-F1, %)",
        "Table 7 Complexity and measured inference efficiency",
        "Table 8 Depth sensitivity of HSTA Block (Macro-F1, %)",
        "Table 9 Cross-protocol transfer results (%)",
        "Table 10 Major confused class pairs of HSTA-Hybrid",
        "Table A1 QUIC label-to-service mapping",
        "Table B1 TLS label-to-service mapping",
    ]
    table_captions = [normalize_spaces(p.text) for p in doc.paragraphs if p.text.strip().startswith("Table ")]
    if table_captions == expected_tables:
        ok(f"{path.name} has continuous Table 1-10 plus Appendix A/B captions")
    else:
        fail(f"{path.name} table captions are {table_captions}, expected {expected_tables}", failures)
    duplicate_table_captions = sorted({caption for caption in table_captions if table_captions.count(caption) > 1})
    if duplicate_table_captions:
        fail(f"{path.name} has duplicate table captions: {duplicate_table_captions}", failures)
    else:
        ok(f"{path.name} has no duplicate table captions")

    document_xml = docx_zip_text(path, r"word/document\.xml")
    table_count = len(doc.tables)
    table_row_count = sum(len(table.rows) for table in doc.tables)
    header_repeat_count = document_xml.count("<w:tblHeader")
    cant_split_count = document_xml.count("<w:cantSplit")
    if header_repeat_count >= table_count:
        ok(f"{path.name} repeats table header rows across page breaks")
    else:
        fail(
            f"{path.name} has {header_repeat_count} repeated table-header markers for {table_count} tables",
            failures,
        )
    if cant_split_count >= table_row_count:
        ok(f"{path.name} prevents table rows from splitting across pages")
    else:
        fail(
            f"{path.name} has {cant_split_count} non-splitting row markers for {table_row_count} table rows",
            failures,
        )

    body_text = "\n".join(p.text for p in doc.paragraphs)
    cjk_chars = sorted({ch for ch in body_text if "\u4e00" <= ch <= "\u9fff"})
    if cjk_chars:
        fail(f"{path.name} contains CJK characters in the English manuscript: {cjk_chars[:12]}", failures)
    else:
        ok(f"{path.name} contains no CJK characters")
    if "[Editable figure placeholder]" in body_text:
        fail(f"{path.name} still contains editable figure placeholder text", failures)
    else:
        ok(f"{path.name} has no editable figure placeholder text")
    repeated_spaces = [match.group(0) for match in re.finditer(r" {2,}", body_text)]
    if repeated_spaces:
        fail(f"{path.name} contains repeated plain spaces in paragraph text", failures)
    else:
        ok(f"{path.name} has no repeated plain spaces in paragraph text")


def check_docx_core_properties(path: Path, doc: Document, failures: list[str], expected_title: str, expected_category: str) -> None:
    props = doc.core_properties
    if props.title == expected_title:
        ok(f"{path.name} core title is controlled")
    else:
        fail(f"{path.name} core title is {props.title!r}, expected {expected_title!r}", failures)
    if props.author == CORE_PROPERTY_AUTHOR and props.last_modified_by == CORE_PROPERTY_AUTHOR:
        ok(f"{path.name} core author metadata is placeholder-controlled")
    else:
        fail(
            f"{path.name} core author metadata is author={props.author!r}, last_modified_by={props.last_modified_by!r}",
            failures,
        )
    if props.keywords == MANUSCRIPT_KEYWORDS:
        ok(f"{path.name} core keywords are controlled")
    else:
        fail(f"{path.name} core keywords are {props.keywords!r}", failures)
    if props.category == expected_category:
        ok(f"{path.name} core category is {expected_category}")
    else:
        fail(f"{path.name} core category is {props.category!r}, expected {expected_category!r}", failures)
    if "python-docx" in " ".join(str(value or "") for value in [props.author, props.comments, props.last_modified_by]):
        fail(f"{path.name} core properties still expose python-docx metadata", failures)
    else:
        ok(f"{path.name} core properties do not expose python-docx metadata")


def check_docx(failures: list[str]) -> None:
    expected = {
        SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_EN.docx": (
            12,
            4,
            [
                "[1,2,35,37]",
                "[4,5,17,27,29,30,32,33,34,35,36,37,39]",
                "[4,5,8,29,30,31,34,36,39]",
                "[3,6,7]",
                "[16,20,24,25,33,34,35,39]",
                "[14]",
                "[26]",
                "[17,32,33,36,37,38]",
                "[28,32,33,38]",
            ],
        ),
        SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_CN.docx": (12, 4, ["[1,2]", "[4,5,8]", "[3,6,7]", "[16,20,24,25]", "[14]", "[26]", "[28]"]),
    }
    for path, (tables, images, citation_markers) in expected.items():
        if not check_exists(path, failures):
            continue
        doc = Document(str(path))
        if len(doc.tables) == tables:
            ok(f"{path.name} has {tables} tables")
        else:
            fail(f"{path.name} table count {len(doc.tables)} != {tables}", failures)
        if len(doc.inline_shapes) == images:
            ok(f"{path.name} has {images} embedded images")
        else:
            fail(f"{path.name} image count {len(doc.inline_shapes)} != {images}", failures)
        body_text = "\n".join(p.text for p in doc.paragraphs)
        for marker in citation_markers:
            if marker in body_text:
                ok(f"{path.name} contains in-text citation marker {marker}")
            else:
                fail(f"{path.name} missing in-text citation marker {marker}", failures)
        if path.name.endswith("_EN.docx"):
            check_english_word_review_format(path, doc, failures)
            check_english_word_structure_hygiene(path, doc, failures)
            check_docx_core_properties(path, doc, failures, MANUSCRIPT_TITLE, "Manuscript")
            figure_captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Figure ")]
            expected_captions = [
                "Figure 1 Schematic of the five-stage HSTA Block",
                "Figure 2 Macro-F1 comparison of main models on four tasks",
                "Figure 3 Attention-position ablation comparison",
                "Figure 4 Cross-protocol transfer comparison",
            ]
            if figure_captions == expected_captions:
                ok(f"{path.name} has continuous Figure 1-4 captions")
            else:
                fail(f"{path.name} figure captions are {figure_captions}, expected {expected_captions}", failures)
        elif path.name.endswith("_CN.docx"):
            check_docx_core_properties(
                path,
                doc,
                failures,
                "面向轻量级加密 QUIC/TLS 流量分类的混合状态空间转换与注意力建模方法",
                "Manuscript",
            )


def check_english_result_tables(failures: list[str]) -> None:
    path = SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_EN.docx"
    if not check_exists(path, failures):
        return
    doc = Document(str(path))
    generator = load_submission_module()
    rows40, rows60 = generator.build_main_result_tables()
    expected_tables = {
        3: rows40,
        4: rows60,
        5: generator.build_scaling_table(),
        6: generator.build_ablation_table(),
        7: generator.build_efficiency_table(),
        8: generator.build_depth_table(),
        9: generator.build_transfer_table(),
    }
    for table_no, expected in expected_tables.items():
        actual = table_rows(doc.tables[table_no - 1])
        if actual == expected:
            ok(f"{path.name} Table {table_no} matches generated results")
        else:
            fail(
                f"{path.name} Table {table_no} differs from generated results; "
                f"actual first rows={actual[:3]}, expected first rows={expected[:3]}",
                failures,
            )


def check_highlights(failures: list[str]) -> None:
    path = SUBMISSION / "Highlights.docx"
    if not check_exists(path, failures):
        return
    doc = Document(str(path))
    check_docx_core_properties(path, doc, failures, "Highlights - " + MANUSCRIPT_TITLE, "Submission support")
    items = [p.text.strip() for p in doc.paragraphs if p.text.strip() and p.text.strip() != "Highlights"]
    if len(items) == 4:
        ok("Highlights.docx contains four bullets")
    else:
        fail(f"Highlights.docx bullet count {len(items)} != 4", failures)
    for item in items:
        if len(item) <= 85:
            ok(f"Highlight <=85 chars: {item}")
        else:
            fail(f"Highlight exceeds 85 chars ({len(item)}): {item}", failures)
    legacy = SUBMISSION / "JNCA_Highlights.docx"
    if legacy.exists():
        fail("Legacy JNCA_Highlights.docx should not exist", failures)
    else:
        ok("Legacy JNCA_Highlights.docx is absent")


def check_statement_docs(failures: list[str]) -> None:
    checks = {
        SUBMISSION / "Author_Statements.docx": [
            "Author contributions (CRediT)",
            "Funding",
            "Declaration of competing interest",
            "Data and code availability",
            "Acknowledgements",
            "placeholder",
        ],
        SUBMISSION / "Declaration_of_Interest_Statement.docx": [
            "Declaration of Interest Statement",
            "no known competing financial interests",
            "placeholder",
        ],
        SUBMISSION / "Cover_Letter.docx": [
            "Journal of Network and Computer Applications",
            "not under consideration elsewhere",
            "competing financial interests",
            "placeholder",
        ],
    }
    for path, needles in checks.items():
        if not check_exists(path, failures):
            continue
        doc = Document(str(path))
        expected_titles = {
            "Author_Statements.docx": "Author Statements - " + MANUSCRIPT_TITLE,
            "Declaration_of_Interest_Statement.docx": "Declaration of Interest - " + MANUSCRIPT_TITLE,
            "Cover_Letter.docx": "Cover Letter - " + MANUSCRIPT_TITLE,
        }
        check_docx_core_properties(path, doc, failures, expected_titles[path.name], "Submission support")
        text = "\n".join(p.text for p in doc.paragraphs)
        for needle in needles:
            if needle in text:
                ok(f"{path.name} contains {needle}")
            else:
                fail(f"{path.name} missing {needle}", failures)


def check_figure_support_files(failures: list[str]) -> None:
    legacy = SUBMISSION / "figures" / "fig5_transfer_macro_f1.png"
    if legacy.exists():
        fail("Legacy fig5_transfer_macro_f1.png should not exist after Figure 1-4 renumbering", failures)
    else:
        ok("Legacy fig5_transfer_macro_f1.png is absent")
    manuscript_figures = {
        "figures/fig1_hsta_architecture.png": (1800, 450),
        "figures/fig2_main_macro_f1.png": (1600, 900),
        "figures/fig3_ablation_macro_f1.png": (1600, 900),
        "figures/fig4_transfer_macro_f1.png": (1600, 900),
    }
    for rel, (min_width, min_height) in manuscript_figures.items():
        fig_path = SUBMISSION / rel
        if check_exists(fig_path, failures):
            if fig_path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n":
                ok(f"{rel} has a PNG header")
            else:
                fail(f"{rel} PNG header check failed", failures)
            check_image_min_size(fig_path, failures, min_width=min_width, min_height=min_height, label=rel)
    if check_exists(FIGURE_CAPTIONS_DOCX, failures):
        doc = Document(str(FIGURE_CAPTIONS_DOCX))
        check_docx_core_properties(path=FIGURE_CAPTIONS_DOCX, doc=doc, failures=failures, expected_title="Figure Captions - " + MANUSCRIPT_TITLE, expected_category="Submission support")
        text = "\n".join(p.text for p in doc.paragraphs)
        for needle in ["Figure 1.", "Figure 2.", "Figure 3.", "Figure 4.", "Graphical Abstract."]:
            if needle in text:
                ok(f"{FIGURE_CAPTIONS_DOCX.name} contains {needle}")
            else:
                fail(f"{FIGURE_CAPTIONS_DOCX.name} missing {needle}", failures)
    if check_exists(GRAPHICAL_ABSTRACT, failures) and GRAPHICAL_ABSTRACT.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n":
        ok("Graphical_Abstract.png has a PNG header")
        check_image_min_size(GRAPHICAL_ABSTRACT, failures, min_width=1328, min_height=531, label="Graphical_Abstract.png")
    elif GRAPHICAL_ABSTRACT.exists():
        fail("Graphical_Abstract.png header check failed", failures)
    if check_exists(GRAPHICAL_ABSTRACT_TIF, failures) and GRAPHICAL_ABSTRACT_TIF.read_bytes()[:4] in {b"II*\x00", b"MM\x00*"}:
        ok("Graphical_Abstract.tif has a TIFF header")
        check_image_min_size(GRAPHICAL_ABSTRACT_TIF, failures, min_width=1328, min_height=531, label="Graphical_Abstract.tif")
    elif GRAPHICAL_ABSTRACT_TIF.exists():
        fail("Graphical_Abstract.tif header check failed", failures)
    png_size = image_size(GRAPHICAL_ABSTRACT) if GRAPHICAL_ABSTRACT.exists() else None
    tif_size = image_size(GRAPHICAL_ABSTRACT_TIF) if GRAPHICAL_ABSTRACT_TIF.exists() else None
    if png_size and tif_size:
        if png_size == tif_size:
            ok("Graphical abstract PNG and TIFF dimensions match")
        else:
            fail(f"Graphical abstract PNG/TIFF dimensions differ: png={png_size}, tif={tif_size}", failures)
    if not check_exists(FIGURE_UPLOAD_ZIP, failures):
        return
    expected = {
        "Figure_1_HSTA_Block.png",
        "Figure_2_Main_Macro_F1.png",
        "Figure_3_Attention_Ablation.png",
        "Figure_4_Cross_Protocol_Transfer.png",
        "Graphical_Abstract.png",
        "Graphical_Abstract.tif",
        "Figure_Captions.docx",
    }
    with zipfile.ZipFile(FIGURE_UPLOAD_ZIP) as zf:
        check_zip_entry_hygiene(FIGURE_UPLOAD_ZIP, zf, failures)
        names = set(zf.namelist())
        missing = expected - names
        if missing:
            fail(f"{FIGURE_UPLOAD_ZIP.name} missing entries: {sorted(missing)}", failures)
        else:
            ok(f"{FIGURE_UPLOAD_ZIP.name} contains figure upload entries")
        expected_sources = {
            "Figure_1_HSTA_Block.png": SUBMISSION / "figures" / "fig1_hsta_architecture.png",
            "Figure_2_Main_Macro_F1.png": SUBMISSION / "figures" / "fig2_main_macro_f1.png",
            "Figure_3_Attention_Ablation.png": SUBMISSION / "figures" / "fig3_ablation_macro_f1.png",
            "Figure_4_Cross_Protocol_Transfer.png": SUBMISSION / "figures" / "fig4_transfer_macro_f1.png",
            "Graphical_Abstract.png": GRAPHICAL_ABSTRACT,
            "Graphical_Abstract.tif": GRAPHICAL_ABSTRACT_TIF,
        }
        for entry, source in expected_sources.items():
            if entry not in names or not source.exists():
                continue
            if zf.read(entry) == source.read_bytes():
                ok(f"{FIGURE_UPLOAD_ZIP.name} entry {entry} matches the generated source file")
            else:
                fail(f"{FIGURE_UPLOAD_ZIP.name} entry {entry} differs from the generated source file", failures)
        for entry in ["Figure_Captions.docx"]:
            source = FIGURE_CAPTIONS_DOCX
            if entry in names and source.exists():
                if zf.read(entry) == source.read_bytes():
                    ok(f"{FIGURE_UPLOAD_ZIP.name} entry {entry} matches the generated source file")
                else:
                    fail(f"{FIGURE_UPLOAD_ZIP.name} entry {entry} differs from the generated source file", failures)
        legacy = [name for name in names if "Figure_5" in name or "fig5" in name]
        if legacy:
            fail(f"{FIGURE_UPLOAD_ZIP.name} contains legacy Figure 5 entries: {legacy}", failures)
        else:
            ok(f"{FIGURE_UPLOAD_ZIP.name} has no legacy Figure 5 entries")


def check_manifest_and_docs(failures: list[str]) -> None:
    checks = {
        SUBMISSION / "submission_upload_manifest.md": [
            "JNCA_HSTA_Encrypted_Traffic_EN.docx",
            "Highlights.docx",
            "Author_Statements.docx",
            "Declaration_of_Interest_Statement.docx",
            "Cover_Letter.docx",
            "Figure_Captions.docx",
            "FIGURES_FOR_UPLOAD.zip",
            "Graphical_Abstract.png",
            "Graphical_Abstract.tif",
            "JNCA_GFA_COMPLIANCE_CHECK.md",
            "FINAL_PLACEHOLDER_REPLACEMENT_CHECK.md",
            "SUBMISSION_IMPORT_FORMAT_CHECK.md",
            "LATEX_IMPORT_TROUBLESHOOTING.md",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "JNCA_recommended_upload_bundle.zip",
            "Final Manual Checks",
        ],
        SUBMISSION / "submission_checklist.md": [
            "Highlights.docx",
            "Author_Statements.docx",
            "Declaration_of_Interest_Statement.docx",
            "Cover_Letter.docx",
            "Figure_Captions.docx",
            "FIGURES_FOR_UPLOAD.zip",
            "Graphical_Abstract.png",
            "Graphical_Abstract.tif",
            "JNCA_GFA_COMPLIANCE_CHECK.md",
            "FINAL_PLACEHOLDER_REPLACEMENT_CHECK.md",
            "SUBMISSION_IMPORT_FORMAT_CHECK.md",
            "LATEX_IMPORT_TROUBLESHOOTING.md",
            "submission_upload_manifest.md",
            "submission_format_audit.md",
            "OVERLEAF_IMPORT_GUIDE.md",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "JNCA_recommended_upload_bundle.zip",
            "PDF header check: ok",
            "zip-entry hygiene",
            "Elsevier `highlights` environment",
            "Word/LaTeX title-keyword-highlight-section alignment",
            "BibTeX integrity",
            "LaTeX package order",
            "source text hygiene",
            "Elsevier template dependency",
            "LaTeX label hygiene",
            "Figure artwork validation",
        ],
        ROOT / "docs" / "jnca_submission_generation.md": [
            "Highlights.docx",
            "Author_Statements.docx",
            "Declaration_of_Interest_Statement.docx",
            "Cover_Letter.docx",
            "Figure_Captions.docx",
            "FIGURES_FOR_UPLOAD.zip",
            "Graphical_Abstract.png",
            "Graphical_Abstract.tif",
            "JNCA_GFA_COMPLIANCE_CHECK.md",
            "FINAL_PLACEHOLDER_REPLACEMENT_CHECK.md",
            "SUBMISSION_IMPORT_FORMAT_CHECK.md",
            "LATEX_IMPORT_TROUBLESHOOTING.md",
            "submission_upload_manifest.md",
            "submission_format_audit.md",
            "OVERLEAF_IMPORT_GUIDE.md",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "JNCA_recommended_upload_bundle.zip",
            "validate_jnca_submission.py",
            "zip-entry hygiene",
            "LaTeX Highlights Policy",
            "LaTeX Package Order Policy",
            "Word Structure Hygiene Policy",
            "Elsevier Frontmatter Policy",
            "Elsevier Template Dependency Policy",
            "Source Text Hygiene Policy",
            "LaTeX Label Hygiene Policy",
            "LaTeX Float Hygiene Policy",
            "Source Package README Policy",
            "Word/LaTeX Alignment Policy",
            "BibTeX Integrity Policy",
            "byte-for-byte",
            "elsarticle.cls not found",
            "first fatal log line",
            "extra top-level folder",
            "Threats to Validity",
        ],
        ROOT / "docs" / "jnca_reference_narrative_notes.md": [
            "JNCA Reference Narrative Notes",
            "SoK: Decoding the Enigma",
            "ET-BERT",
            "NetMamba",
            "TrafficFormer",
            "CESNET-TLS-Year22",
            "DataZoo",
            "DISTILLER",
            "FlowPic",
            "Two-Phase Approach",
            "Classify Traffic Rather Than Flow",
            "SETA++",
            "Rosetta",
            "SAT-Net",
            "data drift",
            "Sweet Danger of Sugar",
            "Narrative lesson",
            "Second-Pass Narrative Lessons",
            "evidence ladder",
            "Negative controls",
            "Manuscript Adjustments Derived From These Papers",
            "Threats to Validity",
            "No training experiments were rerun",
        ],
        SUBMISSION / "submission_format_audit.md": [
            "Word/PDF Structure",
            "12 pt",
            "double spacing",
            "continuous line numbering",
            "page-number field",
            "Separate Submission Files",
            "Declaration_of_Interest_Statement.docx",
            "Recommended direct English Overleaf source",
            "Figure captions",
            "Graphical abstract",
            "Figure artwork integrity",
            "Guide-for-Authors compliance audit",
            "Submission import-format check",
            "LaTeX import troubleshooting",
            "Recommended flat Editorial Manager source",
            "clean ASCII relative entry paths",
            "Word/LaTeX consistency",
            "English Word structure hygiene",
            "BibTeX integrity",
            "package order",
            "frontmatter order",
            "template dependency",
            "source text hygiene",
            "label hygiene",
            "table/figure hygiene",
            "README consistency",
            "template dependency",
            "Remaining Manual Formatting Checks",
        ],
        OVERLEAF_IMPORT_GUIDE: [
            "Why Many Errors Usually Appear",
            "First Fatal Error Map",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_recommended_upload_bundle.zip",
            "main.tex",
            "pdfLaTeX",
            "BibTeX",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "clean ASCII relative paths",
            "elsarticle.cls not found",
            "extra top-level folder",
        ],
        GFA_COMPLIANCE_CHECK: [
            "Abstract word count",
            "Keywords",
            "target: 1-6 keywords",
            "Highlights",
            "Graphical_Abstract.tif",
            "1328 x 531",
            "Do not import `JNCA_recommended_upload_bundle.zip` directly into Overleaf",
        ],
        PLACEHOLDER_REPLACEMENT_CHECK: [
            "Final Placeholder Replacement Check",
            "Required Metadata Replacement",
            "Anonymous Review Decision",
            "Scanned Source Files",
            "Placeholder Scan",
            "Token Inventory",
            "LaTeX Source Placeholder Handling",
            "overleaf_jnca/main.tex",
            "editorial_manager_latex_source/main.tex",
            "Author One",
            "email@example.com",
            "Affiliation 1",
            "Do not upload `JNCA_recommended_upload_bundle.zip` as a direct Overleaf project",
        ],
        IMPORT_FORMAT_CHECK: [
            "Submission Import Format Check",
            "Direct Overleaf Package",
            "Editorial Manager Flat Source Package",
            "Do Not Upload As A LaTeX Project",
            "Error Triage",
            "first fatal line",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "JNCA_recommended_upload_bundle.zip",
            "FIGURES_FOR_UPLOAD.zip",
            "main.tex",
            "pdfLaTeX",
            "BibTeX",
            "XeLaTeX",
            "Zip-entry hygiene",
        ],
        LATEX_TROUBLESHOOTING: [
            "LaTeX Import Troubleshooting",
            "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
            "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
            "JNCA_recommended_upload_bundle.zip",
            "FIGURES_FOR_UPLOAD.zip",
            "main.tex",
            "pdfLaTeX",
            "BibTeX",
            "XeLaTeX",
            "elsarticle.cls not found",
            "elsarticle.dtx",
            "no local TeX engine",
            "first 20 log lines",
            "cascade errors",
        ],
    }
    for path, needles in checks.items():
        if not check_exists(path, failures):
            continue
        text = read_text(path)
        for needle in needles:
            if needle in text:
                ok(f"{path.name} mentions {needle}")
            else:
                fail(f"{path.name} missing {needle}", failures)
    figure_notes = SUBMISSION / "overleaf_jnca" / "figure_notes.md"
    if check_exists(figure_notes, failures):
        text = read_text(figure_notes)
        if "Figures 2, 3, and 4" in text and "Figures 2, 3, and 5" not in text:
            ok("figure_notes.md uses current Figure 1-4 numbering")
        else:
            fail("figure_notes.md has stale result-figure numbering", failures)


def citation_keys(text: str) -> set[str]:
    out: set[str] = set()
    for payload in re.findall(r"\\cite\{([^}]+)\}", text):
        out.update(key.strip() for key in payload.split(",") if key.strip())
    return out


def label_keys(text: str) -> set[str]:
    return set(re.findall(r"\\label\{([^}]+)\}", text))


def ref_keys(text: str) -> set[str]:
    return set(re.findall(r"\\ref\{([^}]+)\}", text))


def latex_label_hygiene_issues(text: str) -> list[str]:
    issues: list[str] = []
    labels = re.findall(r"\\label\{([^}]+)\}", text)
    duplicates = sorted({label for label in labels if labels.count(label) > 1})
    if duplicates:
        issues.append(f"duplicate labels: {duplicates}")
    invalid = sorted(label for label in labels if not re.fullmatch(r"[A-Za-z0-9:-]+", label))
    if invalid:
        issues.append(f"nonportable label names: {invalid}")
    wrong_prefix = sorted(
        label for label in labels if label.startswith(("tab", "fig")) and not label.startswith(("tab:", "fig:"))
    )
    if wrong_prefix:
        issues.append(f"table/figure labels missing standard prefixes: {wrong_prefix}")
    return issues


def main_matter_text(text: str) -> str:
    return text.split(r"\appendix", 1)[0]


def bib_keys(path: Path) -> set[str]:
    if not path.exists():
        return set()
    return set(re.findall(r"@\w+\{([^,]+),", read_text(path)))


def parse_bib_entries(text: str) -> list[tuple[str, str, str]]:
    entries: list[tuple[str, str, str]] = []
    index = 0
    while True:
        match = re.search(r"@(\w+)\s*\{\s*([^,]+)\s*,", text[index:])
        if not match:
            break
        entry_start = index + match.start()
        body_start = index + match.end()
        depth = 1
        cursor = body_start
        while cursor < len(text) and depth:
            if text[cursor] == "{":
                depth += 1
            elif text[cursor] == "}":
                depth -= 1
            cursor += 1
        if depth == 0:
            entries.append((match.group(1).lower(), match.group(2).strip(), text[body_start : cursor - 1]))
        index = cursor
    return entries


def bib_fields(body: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for match in re.finditer(r"(\w+)\s*=\s*\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", body):
        fields[match.group(1).lower()] = normalize_spaces(match.group(2))
    return fields


def check_bib_integrity(path: Path, failures: list[str]) -> dict[str, dict[str, str]]:
    if not check_exists(path, failures):
        return {}
    check_source_text_hygiene(path, failures, ascii_only=True)
    text = read_text(path)
    entries = parse_bib_entries(text)
    if entries:
        ok(f"{path.name} contains {len(entries)} BibTeX entries")
    else:
        fail(f"{path.name} contains no parseable BibTeX entries", failures)
        return {}
    keys = [key for _, key, _ in entries]
    duplicates = sorted({key for key in keys if keys.count(key) > 1})
    if duplicates:
        fail(f"{path.name} has duplicate BibTeX keys: {duplicates}", failures)
    else:
        ok(f"{path.name} has unique BibTeX keys")
    parsed: dict[str, dict[str, str]] = {}
    for entry_type, key, body in entries:
        fields = bib_fields(body)
        parsed[key] = fields | {"entry_type": entry_type}
        required = {"title", "author", "year"}
        if entry_type == "article":
            required.add("journal")
        elif entry_type == "inproceedings":
            required.add("booktitle")
        elif entry_type == "misc":
            if "howpublished" not in fields and not {"eprint", "archiveprefix"} <= fields.keys():
                fail(f"{path.name} entry {key} is misc without howpublished or arXiv-style eprint fields", failures)
        else:
            fail(f"{path.name} entry {key} has unsupported entry type {entry_type!r}", failures)
        missing = sorted(required - fields.keys())
        if missing:
            fail(f"{path.name} entry {key} missing fields: {missing}", failures)
        empty = sorted(name for name, value in fields.items() if not value)
        if empty:
            fail(f"{path.name} entry {key} has empty fields: {empty}", failures)
        year = fields.get("year", "")
        if not re.fullmatch(r"\d{4}", year):
            fail(f"{path.name} entry {key} has invalid year {year!r}", failures)
        combined = " ".join(fields.values()).lower()
        if any(marker in combined for marker in ["placeholder", "todo", "???"]):
            fail(f"{path.name} entry {key} contains placeholder-like text", failures)
    if not any(message.startswith(f"{path.name} entry") for message in failures):
        ok(f"{path.name} BibTeX entries have required fields and valid years")
    return parsed


def unbalanced_envs(text: str) -> list[str]:
    begins = re.findall(r"\\begin\{([^}]+)\}", text)
    ends = re.findall(r"\\end\{([^}]+)\}", text)
    problems: list[str] = []
    for env in sorted(set(begins + ends)):
        if begins.count(env) != ends.count(env):
            problems.append(f"{env}: begin={begins.count(env)} end={ends.count(env)}")
    return problems


def unescaped_special_chars(text: str) -> list[str]:
    problems: list[str] = []
    table_depth = 0
    for lineno, line in enumerate(text.splitlines(), start=1):
        if re.search(r"\\begin\{(?:tabular|tabularx|longtable)\}", line):
            table_depth += 1
        in_table = table_depth > 0
        if line.lstrip().startswith("%"):
            if re.search(r"\\end\{(?:tabular|tabularx|longtable)\}", line):
                table_depth = max(0, table_depth - 1)
            continue
        if not in_table and re.search(r"(?<!\\)&", line):
            problems.append(f"line {lineno}: unescaped &")
        if re.search(r"(?<!\\)#", line):
            problems.append(f"line {lineno}: unescaped #")
        if re.search(r"\\end\{(?:tabular|tabularx|longtable)\}", line):
            table_depth = max(0, table_depth - 1)
    return problems


def unmatched_braces(text: str) -> str | None:
    depth = 0
    escaped = False
    for index, ch in enumerate(text):
        if escaped:
            escaped = False
            continue
        if ch == "\\":
            escaped = True
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth < 0:
                return f"extra closing brace near offset {index}"
    if depth:
        return f"{depth} unmatched opening brace(s)"
    return None


def table_column_count(column_spec: str) -> int:
    cleaned = re.sub(r"@\{[^}]*\}", "", column_spec)
    cleaned = re.sub(r">\{[^}]*\}", "", cleaned)
    cleaned = re.sub(r"\|", "", cleaned)
    count = 0
    index = 0
    while index < len(cleaned):
        ch = cleaned[index]
        if ch in "lcrX":
            count += 1
            index += 1
        elif ch == "p" and index + 1 < len(cleaned) and cleaned[index + 1] == "{":
            count += 1
            index += 2
            depth = 1
            while index < len(cleaned) and depth:
                if cleaned[index] == "{":
                    depth += 1
                elif cleaned[index] == "}":
                    depth -= 1
                index += 1
        else:
            index += 1
    return count


def wide_table_issues(text: str) -> list[str]:
    issues: list[str] = []
    for match in re.finditer(r"\\begin\{table\}.*?\\end\{table\}", text, re.S):
        block = match.group(0)
        if r"\begin{tabularx}" in block:
            continue
        tabular = re.search(r"\\begin\{tabular\}\{([^}]*)\}", block)
        if not tabular:
            continue
        columns = table_column_count(tabular.group(1))
        if columns >= 5 and r"\resizebox{\linewidth}{!}" not in block:
            caption = re.search(r"\\caption\{([^}]*)\}", block)
            label = caption.group(1) if caption else f"offset {match.start()}"
            issues.append(f"{label}: {columns} columns without resizebox")
    return issues


def latex_package_order_issues(text: str) -> list[str]:
    issues: list[str] = []
    if r"\documentclass[preprint,12pt]{elsarticle}" not in text or "ctexart" in text:
        return issues
    required = [r"\usepackage{lineno}", r"\usepackage{hyperref}"]
    missing = [package for package in required if package not in text]
    if missing:
        issues.append(f"missing packages: {missing}")
        return issues
    if text.find(r"\usepackage{lineno}") > text.find(r"\usepackage{hyperref}"):
        issues.append("hyperref should be loaded after lineno in English elsarticle sources")
    return issues


def latex_frontmatter_issues(text: str) -> list[str]:
    issues: list[str] = []
    if r"\documentclass[preprint,12pt]{elsarticle}" not in text or "ctexart" in text:
        return issues
    required_once = [
        r"\begin{document}",
        r"\begin{frontmatter}",
        r"\end{frontmatter}",
        r"\begin{abstract}",
        r"\end{abstract}",
        r"\begin{highlights}",
        r"\end{highlights}",
        r"\begin{keyword}",
        r"\end{keyword}",
        r"\end{document}",
    ]
    for marker in required_once:
        count = text.count(marker)
        if count != 1:
            issues.append(f"{marker} count is {count}, expected 1")
    order = [
        r"\begin{document}",
        r"\begin{frontmatter}",
        r"\title{",
        r"\begin{abstract}",
        r"\end{abstract}",
        r"\begin{highlights}",
        r"\end{highlights}",
        r"\begin{keyword}",
        r"\end{keyword}",
        r"\end{frontmatter}",
        r"\linenumbers",
        r"\section{",
    ]
    positions = [(marker, text.find(marker)) for marker in order]
    missing = [marker for marker, pos in positions if pos < 0]
    if missing:
        issues.append(f"missing frontmatter/order markers: {missing}")
    else:
        for (left_marker, left_pos), (right_marker, right_pos) in zip(positions, positions[1:]):
            if left_pos >= right_pos:
                issues.append(f"{left_marker} should appear before {right_marker}")
    frontmatter = tex_block(text, "frontmatter")
    for env in ["abstract", "highlights", "keyword"]:
        if rf"\begin{{{env}}}" not in frontmatter:
            issues.append(f"{env} environment is not inside frontmatter")
    if r"\affiliation[" in frontmatter:
        issues.append("use legacy-compatible \\address[...] affiliations instead of structured \\affiliation[...]")
    for marker in [r"\address[aff1]{", r"\address[aff2]{"]:
        if marker not in frontmatter:
            issues.append(f"missing legacy-compatible affiliation marker {marker}")
    if r"\maketitle" in text:
        issues.append("elsarticle frontmatter source should not use \\maketitle")
    return issues


def latex_float_hygiene_issues(text: str) -> list[str]:
    issues: list[str] = []
    caption_seen: dict[str, set[str]] = {"figure": set(), "table": set()}
    for match in re.finditer(r"\\begin\{(figure|table)\}(?:\[[^]]*\])?(.*?)\\end\{\1\}", text, re.S):
        env = match.group(1)
        block = match.group(2)
        captions = [normalize_spaces(caption) for caption in re.findall(r"\\caption\{([^{}]*)\}", block)]
        labels = re.findall(r"\\label\{([^}]+)\}", block)
        start_line = text.count("\n", 0, match.start()) + 1
        if len(captions) != 1:
            issues.append(f"{env} near line {start_line} has {len(captions)} captions, expected 1")
        elif not captions[0]:
            issues.append(f"{env} near line {start_line} has an empty caption")
        elif captions[0] in caption_seen[env]:
            issues.append(f"{env} duplicate caption: {captions[0]}")
        elif captions[0].lower() in {"caption", "todo", "placeholder"}:
            issues.append(f"{env} near line {start_line} has placeholder-like caption: {captions[0]}")
        else:
            caption_seen[env].add(captions[0])
        if len(labels) != 1:
            issues.append(f"{env} near line {start_line} has {len(labels)} labels, expected 1")
        elif env == "figure" and not labels[0].startswith("fig:"):
            issues.append(f"figure label {labels[0]!r} should start with fig:")
        elif env == "table" and not labels[0].startswith("tab:"):
            issues.append(f"table label {labels[0]!r} should start with tab:")
        if captions and labels and block.find(r"\label{") < block.find(r"\caption{"):
            issues.append(f"{env} label near line {start_line} appears before its caption")
        if env == "figure" and r"\includegraphics" not in block:
            issues.append(f"figure near line {start_line} has no includegraphics command")
        if env == "table" and not re.search(r"\\begin\{(?:tabular|tabularx|longtable)\}", block):
            issues.append(f"table near line {start_line} has no tabular-like body")
    return issues


def latex_import_path_issues(path: Path, text: str) -> list[str]:
    issues: list[str] = []
    for graphic in re.findall(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", text):
        if "\\" in graphic or graphic.startswith("/") or re.match(r"^[A-Za-z]:", graphic):
            issues.append(f"nonportable graphics path: {graphic}")
        if ".." in Path(graphic.replace("\\", "/")).parts:
            issues.append(f"graphics path escapes the source folder: {graphic}")
    for bib_name in re.findall(r"\\bibliography\{([^}]+)\}", text):
        for name in [item.strip() for item in bib_name.split(",") if item.strip()]:
            if "\\" in name or "/" in name or ".." in name:
                issues.append(f"nonportable bibliography path: {name}")
            if not (path.parent / f"{name}.bib").exists():
                issues.append(f"bibliography file is missing beside {path.name}: {name}.bib")
    for style_name in re.findall(r"\\bibliographystyle\{([^}]+)\}", text):
        if "\\" in style_name or "/" in style_name or ".." in style_name:
            issues.append(f"nonportable bibliography style path: {style_name}")
        if style_name == "elsarticle-num" and not (path.parent / "elsarticle-num.bst").exists():
            issues.append(f"bibliography style file is missing beside {path.name}: elsarticle-num.bst")
    if r"\documentclass[preprint,12pt]{elsarticle}" in text and "ctexart" not in text:
        if text.count(r"\bibliographystyle{elsarticle-num}") != 1:
            issues.append("English elsarticle source should use exactly one \\bibliographystyle{elsarticle-num}")
        if text.count(r"\bibliography{references}") != 1:
            issues.append("English elsarticle source should use exactly one \\bibliography{references}")
    return issues


def check_tex_file(path: Path, failures: list[str], expect_full: bool) -> None:
    if not check_exists(path, failures):
        return
    raw = path.read_bytes()
    text = raw.decode("utf-8")
    ascii_only = r"\documentclass[preprint,12pt]{elsarticle}" in text and "ctexart" not in text
    issues = source_text_hygiene_issues(raw, ascii_only=ascii_only)
    if issues:
        fail(f"{path.name} has source text hygiene issues: {issues}", failures)
    else:
        scope = "ASCII/UTF-8" if ascii_only else "UTF-8"
        ok(f"{path.name} source text hygiene is {scope}-safe")
    problems = unbalanced_envs(text)
    if problems:
        fail(f"{path.name} has unbalanced environments: {problems}", failures)
    else:
        ok(f"{path.name} environments are balanced")
    brace_problem = unmatched_braces(text)
    if brace_problem:
        fail(f"{path.name} has unmatched braces: {brace_problem}", failures)
    else:
        ok(f"{path.name} braces are balanced")
    special_problems = unescaped_special_chars(text)
    if special_problems:
        fail(f"{path.name} has LaTeX special-character issues: {special_problems[:8]}", failures)
    else:
        ok(f"{path.name} has no obvious unescaped & or # issues outside tables")
    table_issues = wide_table_issues(text)
    if table_issues:
        fail(f"{path.name} has wide table formatting issues: {table_issues}", failures)
    else:
        ok(f"{path.name} wide tables use resizebox or tabularx")
    package_issues = latex_package_order_issues(text)
    if package_issues:
        fail(f"{path.name} has LaTeX package-order issues: {package_issues}", failures)
    elif r"\documentclass[preprint,12pt]{elsarticle}" in text and "ctexart" not in text:
        ok(f"{path.name} loads lineno before hyperref")
    frontmatter_issues = latex_frontmatter_issues(text)
    if frontmatter_issues:
        fail(f"{path.name} has Elsevier frontmatter/order issues: {frontmatter_issues}", failures)
    elif r"\documentclass[preprint,12pt]{elsarticle}" in text and "ctexart" not in text:
        ok(f"{path.name} has Elsevier frontmatter in a compile-safe order")
    float_issues = latex_float_hygiene_issues(text)
    if float_issues:
        fail(f"{path.name} has table/figure caption-label issues: {float_issues}", failures)
    else:
        ok(f"{path.name} table/figure floats have captions, labels, and expected bodies")
    import_path_issues = latex_import_path_issues(path, text)
    if import_path_issues:
        fail(f"{path.name} has nonportable or missing import paths: {import_path_issues}", failures)
    elif r"\bibliography" in text or r"\includegraphics" in text:
        ok(f"{path.name} uses portable graphics and bibliography paths")
    if r"\documentclass" in text and path.name == "main.tex" and "ctexart" not in text:
        if r"\documentclass[preprint,12pt]{elsarticle}" in text:
            ok(f"{path.name} uses Elsevier elsarticle preprint class")
        else:
            fail(f"{path.name} does not use the expected Elsevier elsarticle class", failures)
    if r"\documentclass[preprint,12pt]{elsarticle}" in text:
        if r"\ead{email@example.com}" in text:
            ok(f"{path.name} contains Elsevier ead email placeholder")
        else:
            fail(f"{path.name} missing Elsevier ead email placeholder", failures)
        if r"\cortext[cor1]{Corresponding author" in text:
            ok(f"{path.name} contains corresponding-author note")
        else:
            fail(f"{path.name} missing corresponding-author note", failures)
        if r"\address[aff1]{" in text and r"\address[aff2]{" in text and r"\affiliation[" not in text:
            ok(f"{path.name} uses legacy-compatible Elsevier address affiliations")
        else:
            fail(f"{path.name} should use legacy-compatible Elsevier address affiliations", failures)
    graphics = re.findall(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", text)
    missing = [name for name in graphics if not (path.parent / name).exists()]
    if missing:
        fail(f"{path.name} missing graphics: {missing}", failures)
    else:
        ok(f"{path.name} graphics paths exist")
    if "@@" in text:
        fail(f"{path.name} contains internal @@ marker", failures)
    else:
        ok(f"{path.name} has no internal @@ marker")
    labels = label_keys(text)
    refs = ref_keys(text)
    label_issues = latex_label_hygiene_issues(text)
    if label_issues:
        fail(f"{path.name} has LaTeX label hygiene issues: {label_issues}", failures)
    elif labels:
        ok(f"{path.name} labels are unique and portable")
    missing_refs = refs - labels
    if missing_refs:
        fail(f"{path.name} has refs without labels: {sorted(missing_refs)}", failures)
    elif refs:
        ok(f"{path.name} refs resolve to labels")
    if path.parent == BILINGUAL and path.name in {"main.tex", "JNCA_HSTA_Encrypted_Traffic_EN.tex"}:
        main_labels = {label for label in label_keys(main_matter_text(text)) if label.startswith(("tab:", "fig:"))}
        missing_required_refs = main_labels - refs
        if missing_required_refs:
            fail(f"{path.name} does not reference all non-appendix tables/figures: {sorted(missing_required_refs)}", failures)
        else:
            ok(f"{path.name} references all non-appendix tables and figures")
    if expect_full:
        required = [
            "Dataset and task statistics",
            "Adapted advanced baselines",
            "Main results on 40-class",
            "Main results on 60-class",
            "Attention-position ablation",
            "Cross-protocol transfer",
            "Threats to Validity",
            "Author contributions (CRediT)",
        ]
        for needle in required:
            if needle in text:
                ok(f"{path.name} contains {needle}")
            else:
                fail(f"{path.name} missing {needle}", failures)
    if ascii_only:
        ok(f"{path.name} English text is ASCII-safe")
    if r"\documentclass[preprint,12pt]{elsarticle}" in text and "ctexart" not in text:
        highlight_block = re.search(r"\\begin\{highlights\}(.*?)\\end\{highlights\}", text, re.S)
        if highlight_block:
            items = re.findall(r"\\item\s+(.+)", highlight_block.group(1))
            if len(items) == 4:
                ok(f"{path.name} uses Elsevier highlights environment with four items")
            else:
                fail(f"{path.name} highlights item count is {len(items)}, expected 4", failures)
            lengths = [len(item.strip()) for item in items]
            if all(length <= 85 for length in lengths):
                ok(f"{path.name} Highlights are <=85 chars")
            else:
                fail(f"{path.name} Highlights exceed 85 chars: {lengths}", failures)
        else:
            fail(f"{path.name} missing Elsevier highlights environment", failures)
        if r"\section*{Highlights}" in text:
            fail(f"{path.name} still places Highlights as a body section", failures)
        else:
            ok(f"{path.name} does not place Highlights as a body section")
    keyword_block = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", text, re.S)
    if keyword_block:
        keyword_count = len([item for item in keyword_block.group(1).split(r"\sep") if item.strip()])
        if 1 <= keyword_count <= 6:
            ok(f"{path.name} keyword count is {keyword_count}")
        else:
            fail(f"{path.name} keyword count is {keyword_count}, expected 1-6", failures)


def check_latex(failures: list[str]) -> None:
    bib_path = BILINGUAL / "references.bib"
    parsed_bib = check_bib_integrity(bib_path, failures)
    bib = set(parsed_bib)
    submission_bib = SUBMISSION / "overleaf_jnca" / "references.bib"
    if check_exists(submission_bib, failures):
        if normalize_newlines(read_text(submission_bib)) == normalize_newlines(read_text(bib_path)):
            ok("Submission Overleaf references.bib matches bilingual references.bib")
        else:
            fail("Submission Overleaf references.bib differs from bilingual references.bib", failures)
    required_bib = {
        "azab2024traffic",
        "sharma2025encrypted",
        "zhang2023tfegnn",
        "zhao2022mtflowformer",
        "lotfollahi2020deep",
        "zhao2023yatc",
        "meng2022packet",
        "li2023minority",
        "wickramasinghe2025sok",
        "gahtan2024quic",
        "luxemburk2025embedding",
        "dao2024ssm",
        "zhou2025trafficformer",
        "aceto2019mobile",
        "aceto2021distiller",
        "lin2023multimodal",
        "malekghaini2023drift",
        "zhao2025sweet",
        "shapira2021flowpic",
        "wang2023twophase",
        "chen2024multiflow",
        "kattadige2021seta",
        "xie2023rosetta",
        "li2025satnet",
    }
    missing_required_bib = required_bib - bib
    if missing_required_bib:
        fail(f"English BibTeX missing expanded reference keys: {sorted(missing_required_bib)}", failures)
    else:
        ok("English BibTeX contains expanded survey/dataset/model reference coverage")
    for path, expect_full in [
        (BILINGUAL / "main.tex", True),
        (BILINGUAL / "JNCA_HSTA_Encrypted_Traffic_EN.tex", True),
        (SUBMISSION / "overleaf_jnca" / "main.tex", True),
    ]:
        check_tex_file(path, failures, expect_full)
        if path.exists():
            text = read_text(path)
            cites = citation_keys(text)
            missing_cites = cites - bib
            if missing_cites:
                fail(f"{path.name} missing citation keys in bibliography: {sorted(missing_cites)}", failures)
            else:
                ok(f"{path.name} citation keys resolve")
            if "ctexart" not in text:
                required_cites = {
                    "azab2024traffic",
                    "sharma2025encrypted",
                    "gahtan2024quic",
                    "luxemburk2025embedding",
                    "zhao2023yatc",
                    "meng2022packet",
                    "li2023minority",
                    "zhou2025trafficformer",
                    "dao2024ssm",
                    "wang2024netmamba",
                    "wickramasinghe2025sok",
                    "aceto2019mobile",
                    "aceto2021distiller",
                    "lin2023multimodal",
                    "malekghaini2023drift",
                    "zhao2025sweet",
                    "shapira2021flowpic",
                    "wang2023twophase",
                    "chen2024multiflow",
                    "kattadige2021seta",
                    "xie2023rosetta",
                    "li2025satnet",
                }
                missing_required_cites = required_cites - cites
                if missing_required_cites:
                    fail(f"{path.name} missing expanded in-text citation coverage: {sorted(missing_required_cites)}", failures)
                else:
                    ok(f"{path.name} contains expanded in-text citation coverage")
    for value in ["91.09", "90.28", "96.78", "96.29", "30pktTCNET-adapted", "NetMamba-adapted"]:
        text = read_text(BILINGUAL / "main.tex")
        if value in text:
            ok(f"English LaTeX contains {value}")
        else:
            fail(f"English LaTeX missing {value}", failures)


def check_word_latex_alignment(failures: list[str]) -> None:
    docx_path = SUBMISSION / "JNCA_HSTA_Encrypted_Traffic_EN.docx"
    tex_path = SUBMISSION / "overleaf_jnca" / "main.tex"
    if not check_exists(docx_path, failures) or not check_exists(tex_path, failures):
        return
    doc = Document(str(docx_path))
    tex = read_text(tex_path)

    word_title = normalize_spaces(next((p.text for p in doc.paragraphs if p.style.name == "Title" and p.text.strip()), ""))
    latex_title = tex_title(tex)
    if word_title == latex_title == MANUSCRIPT_TITLE:
        ok("Word and LaTeX titles are aligned")
    else:
        fail(f"Word/LaTeX title mismatch: word={word_title!r}, latex={latex_title!r}", failures)

    word_keywords = [normalize_spaces(item) for item in " ".join(extract_section_text(doc, "Keywords")).split(";") if normalize_spaces(item)]
    latex_keywords = tex_keywords(tex)
    if word_keywords == latex_keywords:
        ok("Word and LaTeX keywords are aligned")
    else:
        fail(f"Word/LaTeX keyword mismatch: word={word_keywords}, latex={latex_keywords}", failures)

    word_highlights = [normalize_spaces(item) for item in extract_section_text(doc, "Highlights")]
    latex_highlights = tex_highlights(tex)
    if word_highlights == latex_highlights:
        ok("Word and LaTeX Highlights are aligned")
    else:
        fail(f"Word/LaTeX Highlights mismatch: word={word_highlights}, latex={latex_highlights}", failures)

    word_abstract = normalize_spaces(" ".join(extract_section_text(doc, "Abstract")))
    latex_abstract = tex_plain_text(tex_block(tex, "abstract"))
    for marker in [
        "TLS 1.3",
        "QUIC",
        "HSTA Block",
        "QUIC-40",
        "QUIC-60",
        "TLS-40",
        "TLS-60",
        "91.09",
        "90.28",
        "96.78",
        "96.29",
    ]:
        if marker in word_abstract and marker in latex_abstract:
            ok(f"Word and LaTeX abstracts both contain {marker}")
        else:
            fail(f"Word/LaTeX abstract marker mismatch for {marker}", failures)

    word_headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading") and re.match(r"^\d+ ", p.text.strip())]
    latex_sections = re.findall(r"\\section\{([^}]+)\}", tex)
    expected_sections = [
        "Introduction",
        "Background and Related Work",
        "Problem Definition and Motivation",
        "Proposed Method",
        "Experimental Setup",
        "Results and Analysis",
        "Discussion",
        "Threats to Validity",
        "Conclusion",
    ]
    word_section_titles = [re.sub(r"^\d+\s+", "", heading) for heading in word_headings[: len(expected_sections)]]
    latex_main_sections = [title for title in latex_sections if not title.startswith("Appendix")][: len(expected_sections)]
    if word_section_titles == latex_main_sections == expected_sections:
        ok("Word and LaTeX main section titles are aligned")
    else:
        fail(
            f"Word/LaTeX section mismatch: word={word_section_titles}, latex={latex_main_sections}",
            failures,
        )


def check_zip(failures: list[str]) -> None:
    english_zip_entries = {
        "main.tex",
        "JNCA_HSTA_Encrypted_Traffic_EN.tex",
        "references.bib",
        "elsarticle-num.bst",
        "README.md",
        "latexmkrc",
        "figures/fig1_hsta_architecture.png",
        "figures/fig2_main_macro_f1.png",
        "figures/fig3_ablation_macro_f1.png",
        "figures/fig4_transfer_macro_f1.png",
    }
    required = {
        ENGLISH_OVERLEAF_ZIP: english_zip_entries,
        OVERLEAF_DIRECT_ZIP: english_zip_entries,
        SUBMISSION / "overleaf_jnca_package.zip": english_zip_entries | {"README_overleaf.md", "figure_notes.md", "FIGURE_BUILD_NOTE.txt"},
        EDITORIAL_SOURCE_ZIP: {
            "main.tex",
            "JNCA_HSTA_Encrypted_Traffic_EN.tex",
            "references.bib",
            "elsarticle-num.bst",
            "README.md",
            "latexmkrc",
            "fig1_hsta_architecture.png",
            "fig2_main_macro_f1.png",
            "fig3_ablation_macro_f1.png",
            "fig4_transfer_macro_f1.png",
        },
        ROOT / "paper" / "jnca_bilingual_latex_package.zip": {
            "main.tex",
            "main_cn.tex",
            "JNCA_HSTA_Encrypted_Traffic_EN.tex",
            "JNCA_HSTA_Encrypted_Traffic_CN.tex",
            "references.bib",
            "elsarticle-num.bst",
            "README.md",
            "latexmkrc",
            "figures/fig1_hsta_architecture.png",
        },
        ROOT / "paper" / "jnca_chinese_xelatex_package.zip": {
            "main.tex",
            "JNCA_HSTA_Encrypted_Traffic_CN.tex",
            "README.md",
            "latexmkrc",
            "figures/fig1_hsta_architecture.png",
        },
    }
    for path, expected in required.items():
        if not check_exists(path, failures):
            continue
        with zipfile.ZipFile(path) as zf:
            check_zip_entry_hygiene(path, zf, failures)
            names = set(zf.namelist())
            check_zip_readme_consistency(path, zf, names, failures)
            if all("/" in name for name in names if name and not name.endswith("/")):
                fail(f"{path.name} appears to contain a top-level folder instead of root-level files", failures)
            else:
                ok(f"{path.name} has root-level files for Overleaf import")
            missing = expected - names
            if missing:
                fail(f"{path.name} missing entries: {sorted(missing)}", failures)
            else:
                ok(f"{path.name} contains required entries")
            if "latexmkrc" in names:
                latexmk = zf.read("latexmkrc").decode("utf-8").strip()
                if path.name == "jnca_chinese_xelatex_package.zip":
                    expected_mode = "$pdf_mode = 5;"
                else:
                    expected_mode = "$pdf_mode = 1;"
                if latexmk == expected_mode:
                    ok(f"{path.name} has {expected_mode}")
                else:
                    fail(f"{path.name} latexmkrc is {latexmk!r}, expected {expected_mode!r}", failures)
            if "elsarticle-num.bst" in names:
                if check_exists(OFFICIAL_ELSARTICLE_BST, failures):
                    packaged_bst = zf.read("elsarticle-num.bst")
                    official_bst = OFFICIAL_ELSARTICLE_BST.read_bytes()
                    if packaged_bst == official_bst:
                        ok(f"{path.name} elsarticle-num.bst matches the official template copy")
                    else:
                        fail(f"{path.name} elsarticle-num.bst differs from the official template copy", failures)
            if "references.bib" in names and (SUBMISSION / "overleaf_jnca" / "references.bib").exists():
                zipped_bib_bytes = zf.read("references.bib")
                bib_issues = source_text_hygiene_issues(zipped_bib_bytes, ascii_only=True)
                if bib_issues:
                    fail(f"{path.name} references.bib has source text hygiene issues: {bib_issues}", failures)
                else:
                    ok(f"{path.name} references.bib source text hygiene is ASCII/UTF-8-safe")
                zipped_bib = zipped_bib_bytes.decode("utf-8")
                canonical_bib = read_text(SUBMISSION / "overleaf_jnca" / "references.bib")
                if normalize_newlines(zipped_bib) == normalize_newlines(canonical_bib):
                    ok(f"{path.name} references.bib matches the submission Overleaf copy")
                else:
                    fail(f"{path.name} references.bib differs from the submission Overleaf copy", failures)
            if path.name in {"jnca_english_overleaf_package.zip", "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip", "overleaf_jnca_package.zip", "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip"} and any("CN.tex" in name or "main_cn" in name for name in names):
                fail(f"{path.name} contains Chinese TeX files", failures)
            elif path.name in {"jnca_english_overleaf_package.zip", "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip", "overleaf_jnca_package.zip", "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip"}:
                ok(f"{path.name} excludes Chinese TeX files")
            if "main.tex" in names:
                main_bytes = zf.read("main.tex")
                main_text = main_bytes.decode("utf-8")
                if path.name in {"jnca_english_overleaf_package.zip", "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip", "overleaf_jnca_package.zip", "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip"}:
                    main_issues = source_text_hygiene_issues(main_bytes, ascii_only=True)
                    if main_issues:
                        fail(f"{path.name} main.tex has source text hygiene issues: {main_issues}", failures)
                    else:
                        ok(f"{path.name} main.tex source text hygiene is ASCII/UTF-8-safe")
                    forbidden = [
                        name
                        for name in names
                        if Path(name).suffix.lower() in {".docx", ".pdf", ".zip", ".tif", ".tiff"}
                    ]
                    if forbidden:
                        fail(f"{path.name} contains non-source upload files: {sorted(forbidden)}", failures)
                    else:
                        ok(f"{path.name} excludes Word/PDF/nested-zip/artwork-only files")
                    if r"\documentclass[preprint,12pt]{elsarticle}" in main_text and "ctexart" not in main_text:
                        ok(f"{path.name} main.tex is English elsarticle")
                    else:
                        fail(f"{path.name} main.tex is not the expected English elsarticle file", failures)
                    if "pdfLaTeX" in zf.read("README.md").decode("utf-8"):
                        ok(f"{path.name} README mentions pdfLaTeX")
                    else:
                        fail(f"{path.name} README does not mention pdfLaTeX", failures)
                    graphics = re.findall(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", main_text)
                    missing_graphics = [graphic for graphic in graphics if graphic not in names]
                    if missing_graphics:
                        fail(f"{path.name} main.tex references graphics missing from zip: {missing_graphics}", failures)
                    else:
                        ok(f"{path.name} main.tex graphics are present inside the zip")
                    if path.name in {"jnca_english_overleaf_package.zip", "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip", "overleaf_jnca_package.zip"}:
                        if graphics and all(graphic.startswith("figures/") for graphic in graphics):
                            ok(f"{path.name} uses Overleaf-style figures/ graphics paths")
                        else:
                            fail(f"{path.name} does not consistently use figures/ graphics paths: {graphics}", failures)
                    if path.name == "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip":
                        nested_entries = [name for name in names if "/" in name.rstrip("/")]
                        if nested_entries:
                            fail(f"{path.name} contains subfolder entries that Editorial Manager may reject: {nested_entries[:8]}", failures)
                        else:
                            ok(f"{path.name} is flat for Editorial Manager import")
                        if any("/" in graphic or "\\" in graphic for graphic in graphics):
                            fail(f"{path.name} main.tex still references subfolder graphics: {graphics}", failures)
                        elif all(graphic in names for graphic in graphics):
                            ok(f"{path.name} main.tex uses root-level graphics")
                        else:
                            fail(f"{path.name} main.tex references missing root-level graphics: {graphics}", failures)
    if ENGLISH_OVERLEAF_ZIP.exists() and OVERLEAF_DIRECT_ZIP.exists():
        with zipfile.ZipFile(ENGLISH_OVERLEAF_ZIP) as source_zip, zipfile.ZipFile(OVERLEAF_DIRECT_ZIP) as direct_zip:
            source_names = set(source_zip.namelist())
            direct_names = set(direct_zip.namelist())
            if source_names != direct_names:
                fail("Direct Overleaf zip entries differ from the canonical English Overleaf zip", failures)
            else:
                mismatched = [
                    name
                    for name in sorted(source_names)
                    if not name.endswith("/") and source_zip.read(name) != direct_zip.read(name)
                ]
                if mismatched:
                    fail(f"Direct Overleaf zip content differs from canonical package: {mismatched[:5]}", failures)
                else:
                    ok("Direct Overleaf zip matches the canonical English Overleaf package")
    folder_main = SUBMISSION / "overleaf_jnca" / "main.tex"
    if folder_main.exists() and OVERLEAF_DIRECT_ZIP.exists():
        with zipfile.ZipFile(OVERLEAF_DIRECT_ZIP) as zf:
            direct_main = zf.read("main.tex").decode("utf-8")
        if normalize_newlines(folder_main.read_text(encoding="utf-8")) == normalize_newlines(direct_main):
            ok("Submission overleaf_jnca/main.tex matches the direct Overleaf zip")
        else:
            fail("Submission overleaf_jnca/main.tex differs from the direct Overleaf zip", failures)


def check_upload_bundle(failures: list[str]) -> None:
    expected = {
        "JNCA_HSTA_Encrypted_Traffic_EN.docx",
        "JNCA_HSTA_Encrypted_Traffic_EN.pdf",
        "Highlights.docx",
        "Author_Statements.docx",
        "Declaration_of_Interest_Statement.docx",
        "Cover_Letter.docx",
        "Figure_Captions.docx",
        "Graphical_Abstract.png",
        "Graphical_Abstract.tif",
        "FIGURES_FOR_UPLOAD.zip",
        "submission_upload_manifest.md",
        "submission_checklist.md",
        "submission_format_audit.md",
        "JNCA_GFA_COMPLIANCE_CHECK.md",
        "FINAL_PLACEHOLDER_REPLACEMENT_CHECK.md",
        "SUBMISSION_IMPORT_FORMAT_CHECK.md",
        "LATEX_IMPORT_TROUBLESHOOTING.md",
        "OVERLEAF_IMPORT_GUIDE.md",
        "UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip",
        "JNCA_EDITORIAL_MANAGER_LATEX_SOURCE.zip",
    }
    if not check_exists(UPLOAD_BUNDLE, failures):
        return
    with zipfile.ZipFile(UPLOAD_BUNDLE) as zf:
        check_zip_entry_hygiene(UPLOAD_BUNDLE, zf, failures)
        names = set(zf.namelist())
        missing = expected - names
        if missing:
            fail(f"{UPLOAD_BUNDLE.name} missing entries: {sorted(missing)}", failures)
        else:
            ok(f"{UPLOAD_BUNDLE.name} contains recommended upload entries")
        unexpected = [name for name in names if "CN" in name or "chinese" in name.lower() or "bilingual" in name.lower()]
        if unexpected:
            fail(f"{UPLOAD_BUNDLE.name} contains optional/non-English files: {unexpected}", failures)
        else:
            ok(f"{UPLOAD_BUNDLE.name} excludes Chinese and bilingual support files")
        if "main.tex" in names:
            fail(f"{UPLOAD_BUNDLE.name} contains root main.tex; it should remain a handoff bundle, not a direct Overleaf source zip", failures)
        else:
            ok(f"{UPLOAD_BUNDLE.name} is not shaped like a direct Overleaf source zip")
        nested = zf.read("UPLOAD_TO_OVERLEAF_JNCA_ENGLISH.zip")
        with zipfile.ZipFile(io.BytesIO(nested)) as nested_zip:
            nested_issues = zip_entry_hygiene_issues(nested_zip.namelist())
            if nested_issues:
                fail(f"Nested English Overleaf zip has unsafe or nonportable entries: {nested_issues[:8]}", failures)
            else:
                ok("Nested English Overleaf zip has clean entry paths")
            nested_names = set(nested_zip.namelist())
            nested_required = {"main.tex", "references.bib", "elsarticle-num.bst", "latexmkrc", "README.md"}
            nested_missing = nested_required - nested_names
            if nested_missing:
                fail(f"Nested English Overleaf zip missing entries: {sorted(nested_missing)}", failures)
            else:
                ok("Nested English Overleaf zip contains core entries")
            nested_main = nested_zip.read("main.tex").decode("utf-8")
            if r"\documentclass[preprint,12pt]{elsarticle}" in nested_main and "ctexart" not in nested_main:
                ok("Nested English Overleaf main.tex is elsarticle/pdfLaTeX-oriented")
            else:
                fail("Nested English Overleaf main.tex is not the expected English elsarticle file", failures)
    if ENGLISH_OVERLEAF_ZIP.exists() and UPLOAD_BUNDLE.stat().st_mtime < ENGLISH_OVERLEAF_ZIP.stat().st_mtime:
        fail("Recommended upload bundle is older than the English Overleaf zip; regenerate submission after LaTeX package", failures)
    else:
        ok("Recommended upload bundle is not older than the English Overleaf zip")


def main() -> int:
    failures: list[str] = []
    check_pdf(failures)
    check_docx(failures)
    check_english_result_tables(failures)
    check_highlights(failures)
    check_statement_docs(failures)
    check_figure_support_files(failures)
    check_manifest_and_docs(failures)
    check_latex(failures)
    check_word_latex_alignment(failures)
    check_zip(failures)
    check_upload_bundle(failures)
    if failures:
        print(f"\nValidation failed with {len(failures)} issue(s).")
        return 1
    print("\nValidation passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
