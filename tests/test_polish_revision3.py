from __future__ import annotations

import hashlib
import os
import re
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from scripts.generate_hsta_drawio_figures import TRANSFER_TASKS, load_figure_data


ROOT = Path(os.environ.get("HSTA_PROJECT_ROOT", Path(__file__).resolve().parents[1])).resolve()
ARCHIVE = ROOT / "paper" / "源稿归档"
SOURCE = ARCHIVE / "加密QUIC_TLS流量分类_修改版2.docx"
TARGET = ARCHIVE / "加密QUIC_TLS流量分类_修改版3.docx"
FIGURE_DIR = ROOT / "paper" / "figures_drawio"


class Revision3Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.assertTrue = unittest.TestCase().assertTrue
        cls.assertTrue(TARGET.exists(), str(TARGET))
        cls.doc = Document(TARGET)
        cls.text = "\n".join(paragraph.text for paragraph in cls.doc.paragraphs)
        cls.figure_data = load_figure_data(ROOT)

    def test_source_is_preserved_and_target_is_complete(self):
        self.assertTrue(SOURCE.exists())
        self.assertEqual(len(self.doc.tables), 11)
        self.assertEqual(len(self.doc.inline_shapes), 7)
        with zipfile.ZipFile(TARGET) as archive:
            document_xml = archive.read("word/document.xml")
            self.assertEqual(document_xml.count(b"<m:oMath"), 23)

    def test_strengthened_structure_and_claims(self):
        self.assertEqual(
            self.doc.paragraphs[0].text,
            "HSTA：面向轻量级加密 QUIC/TLS 流量分类的混合状态空间转换与注意力模型",
        )
        self.assertEqual(
            self.doc.paragraphs[1].text,
            "HSTA: Hybrid State-Space Transformation and Attention for Lightweight Encrypted QUIC/TLS Traffic Classification",
        )
        self.assertIn("2.6 研究空白与本文定位（Research Gap and Positioning）", self.text)
        method_headings = [
            paragraph.text
            for paragraph in self.doc.paragraphs
            if paragraph.style.name.startswith("Heading") and paragraph.text.startswith("3.")
        ]
        self.assertEqual(
            method_headings,
            [
                "3.1 问题定义（Problem Formulation）",
                "3.2 设计动机（Design Motivation）",
                "3.3 总体架构（Overall Architecture）",
                "3.4 状态空间序列编码（State-Space Sequence Encoding）",
                "3.5 表示转换与高层注意力（Representation Transformation and High-Level Attention）",
                "3.6 训练配置（Training Configuration）",
                "3.7 模型规模与部署特性（Model Scale and Deployment Characteristics）",
            ],
        )
        self.assertIn("6 适用范围与拓展方向（Scope and Future Extensions）", self.text)
        self.assertNotIn("6 局限性（Limitations）", self.text)
        self.assertNotIn("采用审慎表述", self.text)
        self.assertNotIn("不声称 HSTA", self.text)
        self.assertNotIn("严格复现", self.text)
        for value in (
            "91.09±0.20%",
            "90.28±0.39%",
            "96.78±0.14%",
            "96.29±0.17%",
            "41.5%",
            "65.9%",
        ):
            self.assertIn(value, self.text)
        self.assertNotRegex(self.text, r"\[\d+(?:,\d+|[-–]\d+)+\]")

    def test_figures_are_current_and_numbered_continuously(self):
        captions = [
            paragraph.text.strip()
            for paragraph in self.doc.paragraphs
            if re.match(r"^图\d+\s", paragraph.text.strip())
        ]
        self.assertEqual([re.match(r"^图(\d+)", caption).group(1) for caption in captions], list("1234567"))

        expected_hashes = {
            hashlib.sha256((FIGURE_DIR / filename).read_bytes()).hexdigest()
            for filename in (
                "fig1_hsta_architecture.png",
                "fig2_main_macro_f1.png",
                "fig3_attention_position_ablation.png",
                "fig4_parameter_performance.png",
                "fig5_cross_protocol_transfer.png",
                "fig6_confusion_patterns.png",
                "fig7_data_processing_pipeline.png",
            )
        }
        with zipfile.ZipFile(TARGET) as archive:
            actual_hashes = {
                hashlib.sha256(archive.read(name)).hexdigest()
                for name in archive.namelist()
                if name.startswith("word/media/")
            }
        self.assertTrue(expected_hashes.issubset(actual_hashes))

    def test_tables_remain_english(self):
        table_text = "\n".join(
            cell.text
            for table in self.doc.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertNotRegex(table_text, r"[\u3400-\u4dbf\u4e00-\u9fff]")
        for table in self.doc.tables:
            header_properties = table.rows[0]._tr.find(qn("w:trPr"))
            self.assertIsNotNone(header_properties)
            self.assertIsNotNone(header_properties.find(qn("w:tblHeader")))
            for row in table.rows:
                row_properties = row._tr.find(qn("w:trPr"))
                self.assertIsNotNone(row_properties.find(qn("w:cantSplit")))

    def test_result_tables_match_validated_csv_data(self):
        def metric(value: tuple[float, float | None]) -> str:
            mean, std = value
            return f"{mean:.2f}" if std is None else f"{mean:.2f}±{std:.2f}"

        table2_hsta = [cell.text for cell in self.doc.tables[1].rows[-1].cells]
        self.assertEqual(table2_hsta[2], metric(self.figure_data.main["HSTA"]["QUIC-40"]))
        self.assertEqual(table2_hsta[4], metric(self.figure_data.main["HSTA"]["TLS-40"]))

        table3_hsta = [cell.text for cell in self.doc.tables[2].rows[-1].cells]
        self.assertEqual(table3_hsta[2], metric(self.figure_data.main["HSTA"]["QUIC-60"]))
        self.assertEqual(table3_hsta[4], metric(self.figure_data.main["HSTA"]["TLS-60"]))

        ablation_labels = {
            "No attention": "No attention",
            "Front attention": "Front attention",
            "Middle attention": "Middle attention",
            "HSTA (full model)": "HSTA",
        }
        for row in self.doc.tables[4].rows[1:]:
            cells = [cell.text for cell in row.cells]
            source_label = ablation_labels[cells[0]]
            for column, task in enumerate(("QUIC-40", "QUIC-60", "TLS-40", "TLS-60"), start=1):
                self.assertEqual(cells[column], metric(self.figure_data.ablation[source_label][task]))

        transfer_table = self.doc.tables[7]
        for row, task in zip(transfer_table.rows[1:], TRANSFER_TASKS):
            cells = [cell.text for cell in row.cells]
            self.assertEqual(float(cells[2]), self.figure_data.transfer["Zero-shot"][task])
            self.assertEqual(float(cells[3]), self.figure_data.transfer["LoRA"][task])
            self.assertEqual(float(cells[5]), self.figure_data.transfer["Full fine-tuning"][task])
            self.assertEqual(float(cells[6]), self.figure_data.transfer["Target scratch"][task])

    def test_references_are_complete_and_citations_are_valid(self):
        references = [paragraph.text for paragraph in self.doc.paragraphs if re.match(r"^\[\d+\]", paragraph.text)]
        self.assertEqual(len(references), 28)
        reference_ids = {int(re.match(r"^\[(\d+)\]", value).group(1)) for value in references}
        self.assertEqual(reference_ids, set(range(1, 29)))
        cited_ids = {int(value) for value in re.findall(r"\[(\d+)\]", self.text)}
        self.assertTrue(cited_ids.issubset(reference_ids))


if __name__ == "__main__":
    unittest.main()
